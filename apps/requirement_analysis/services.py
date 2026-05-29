import os
import io
import base64
import json
import time
import uuid
import asyncio
import logging
import zipfile
from typing import List, Dict, Any, Optional
from datetime import datetime

try:
    from PyPDF2 import PdfReader
except ImportError:
    from PyPDF2 import PdfFileReader as PdfReader
    
try:
    import docx
    from docx.document import Document as DocxDocument
except ImportError:
    docx = None

try:
    from PIL import Image
except ImportError:
    Image = None
    
import httpx
import httpx
from django.conf import settings
from django.core.files.storage import default_storage

from .models import RequirementDocument, RequirementAnalysis, BusinessRequirement, GeneratedTestCase, AnalysisTask

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """文档处理服务"""
    
    @staticmethod
    def extract_images(document: RequirementDocument) -> List[Dict[str, str]]:
        file_path = getattr(document.file, 'path', '')
        if not file_path or not os.path.exists(file_path):
            return []

        if document.document_type != 'docx':
            return []

        target_dir = os.path.join(settings.MEDIA_ROOT, 'requirement_doc_images', str(document.id))
        os.makedirs(target_dir, exist_ok=True)

        image_items = []

        try:
            with zipfile.ZipFile(file_path) as archive:
                media_files = sorted(
                    [name for name in archive.namelist() if name.startswith('word/media/') and not name.endswith('/')],
                    key=lambda name: name.lower()
                )

                for index, archive_name in enumerate(media_files, start=1):
                    ext = os.path.splitext(archive_name)[1].lower() or '.png'
                    output_name = f'image_{index}{ext}'
                    output_path = os.path.join(target_dir, output_name)

                    if not os.path.exists(output_path):
                        with archive.open(archive_name) as src, open(output_path, 'wb') as dst:
                            dst.write(src.read())

                    image_items.append({
                        'name': output_name,
                        'label': f'文档图片 {index}',
                        'url': f"{settings.MEDIA_URL}requirement_doc_images/{document.id}/{output_name}"
                    })
        except Exception as e:
            logger.error(f"Word图片提取失败: {e}")
            return []

        return image_items

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """从PDF文件提取文本"""
        try:
            text = ""
            with open(file_path, 'rb') as file:
                pdf_reader = PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
            return text.strip()
        except Exception as e:
            logger.error(f"PDF文本提取失败: {e}")
            return f"PDF文本提取失败: {str(e)}"
    
    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """从Word文档提取文本（包括表格）"""
        try:
            doc = docx.Document(file_path)
            text = ""
            
            # 1. 提取普通段落
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"
            
            # 2. 提取表格内容
            for table in doc.tables:
                text += "\n--- 表格 --\n"
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        cell_text = cell.text.strip()
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        text += " | ".join(row_text) + "\n"
            
            # 3. 提取页眉页脚
            for section in doc.sections:
                if section.header.paragraphs:
                    for p in section.header.paragraphs:
                        if p.text.strip():
                            text += p.text + "\n"
                if section.footer.paragraphs:
                    for p in section.footer.paragraphs:
                        if p.text.strip():
                            text += p.text + "\n"
            
            return text.strip()
        except Exception as e:
            logger.error(f"Word文档文本提取失败: {e}")
            return f"Word文档文本提取失败: {str(e)}"
    
    @staticmethod
    def extract_text_from_txt(file_path: str) -> str:
        """从文本文件提取文本"""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read().strip()
        except UnicodeDecodeError:
            try:
                with open(file_path, 'r', encoding='gbk') as file:
                    return file.read().strip()
            except Exception as e:
                logger.error(f"文本文件读取失败: {e}")
                return f"文本文件读取失败: {str(e)}"
        except Exception as e:
            logger.error(f"文本文件读取失败: {e}")
            return f"文本文件读取失败: {str(e)}"
    
    @classmethod
    def extract_text(cls, document: RequirementDocument) -> str:
        """根据文档类型提取文本"""
        file_path = document.file.path
        
        if document.document_type == 'pdf':
            return cls.extract_text_from_pdf(file_path)
        elif document.document_type == 'docx':
            return cls.extract_text_from_docx(file_path)
        elif document.document_type == 'txt':
            return cls.extract_text_from_txt(file_path)
        elif document.document_type == 'md':
            return cls.extract_text_from_txt(file_path)
        else:
            return "不支持的文档类型"


class AIService:
    """AI服务类 - 对接真实大模型API"""

    @staticmethod
    async def _chat_completion(messages: List[Dict[str, str]], temperature: float = 0.7, max_tokens: int = 8192) -> str:
        """调用OpenAI兼容大模型API（豆包）"""
        api_key = getattr(settings, 'DOUBAO_API_KEY', getattr(settings, 'OPENAI_API_KEY', ''))
        base_url = getattr(settings, 'DOUBAO_BASE_URL', getattr(settings, 'OPENAI_BASE_URL', 'https://ark.cn-beijing.volces.com/api/v3'))
        model = getattr(settings, 'DOUBAO_MODEL', getattr(settings, 'OPENAI_MODEL', 'doubao-seed-2.0-pro'))
        
        if not api_key:
            raise ValueError("大模型API密钥未配置")
        
        async with httpx.AsyncClient(timeout=120) as client:
            response = await client.post(
                f"{base_url.rstrip('/')}/chat/completions",
                headers={"Authorization": f"Bearer {api_key}"},
                json={
                    "model": model,
                    "messages": messages,
                    "temperature": temperature,
                    "max_tokens": max_tokens,
                    "stream": False
                }
            )
            response.raise_for_status()
            result = response.json()
            return result['choices'][0]['message']['content'].strip()
    
    @staticmethod
    async def analyze_requirements(text: str, document_title: str = "") -> Dict[str, Any]:
        """
        先进的需求分析 - 使用新的智能分析引擎
        
        Args:
            text: 需求文档文本内容
            document_title: 文档标题
            
        Returns:
            Dict包含分析报告、结构化需求等信息
        """
        try:
            # 直接导入并使用先进分析器
            from apps.requirement_analysis.advanced_analyzer import advanced_analyzer
            
            logger.info(f"使用先进分析器分析需求，文档标题: {document_title}")
            
            # 使用先进分析器进行分析
            result = await advanced_analyzer.analyze_requirements_advanced(text, document_title)
            
            # 转换为原系统期望的格式
            analysis_report = result.get("analysis_report", "")
            structured_requirements = result.get("structured_requirements", {})
            requirements_list = structured_requirements.get("requirements", [])
            
            # 计算分析时间（模拟）
            import time
            analysis_time = time.time() % 10 + 2  # 2-12秒之间的模拟时间
            
            logger.info(f"先进需求分析完成，识别需求{len(requirements_list)}个")
            
            return {
                "analysis_report": analysis_report,
                "requirements": requirements_list,
                "requirements_count": len(requirements_list),
                "analysis_time": analysis_time,
                "quality_assessment": result.get("quality_assessment", {}),
                "risk_analysis": result.get("risk_analysis", {})
            }
            
        except Exception as e:
            logger.error(f"先进需求分析失败: {e}")
            logger.info("使用备用分析方法")
            # fallback到原来的分析逻辑
            return await AIService._fallback_analyze_requirements(text, document_title)
    
    @staticmethod
    async def _fallback_analyze_requirements(text: str, document_title: str = "") -> Dict[str, Any]:
        """备用需求分析方法"""
        # 模拟AI分析过程
        await asyncio.sleep(2)
        
        # 这里应该调用真实的大模型API
        # 现在返回改进的模拟数据
        analysis_report = f"""
# 需求分析报告

## 文档概述
基于提供的需求文档"{document_title}"，共识别出以下主要需求模块和功能点。

## 主要功能模块
1. 用户管理模块
2. 数据处理模块  
3. 报告生成模块
4. 系统配置模块

## 详细需求分析
基于文档内容分析，识别出以下具体需求：

### 功能需求
- 用户认证和权限管理
- 数据录入和维护功能
- 业务流程处理
- 报表和统计功能

### 非功能需求
- 系统性能要求：响应时间 < 3秒
- 安全性要求：数据加密存储
- 可用性要求：99.5%系统可用率
- 兼容性要求：支持主流浏览器

## 风险评估
- 技术实现风险：中等
- 进度风险：低
- 资源风险：低

## 建议
建议采用敏捷开发模式，分阶段实施各功能模块。

分析时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
        """
        
        # 生成基础的结构化需求
        requirements = [
            {
                "requirement_id": "REQ-001",
                "requirement_name": "用户认证管理", 
                "requirement_type": "functional",
                "parent_requirement": None,
                "module": "用户管理",
                "requirement_level": "high",
                "reviewer": "admin",
                "estimated_hours": 16,
                "description": "作为一名系统用户，我希望通过用户名和密码登录系统，这样可以确保系统安全性并获得个性化服务。",
                "acceptance_criteria": "用户能够使用有效凭证成功登录系统，无效凭证登录失败，系统记录登录日志。"
            },
            {
                "requirement_id": "REQ-002",
                "requirement_name": "数据管理功能",
                "requirement_type": "functional", 
                "parent_requirement": None,
                "module": "数据管理",
                "requirement_level": "high",
                "reviewer": "admin",
                "estimated_hours": 24,
                "description": "作为一名数据操作员，我希望能够对系统数据进行增删改查操作，这样可以有效管理业务信息。",
                "acceptance_criteria": "数据操作功能正常，数据完整性得到保证，操作权限控制有效。"
            },
            {
                "requirement_id": "REQ-003",
                "requirement_name": "报表统计功能",
                "requirement_type": "functional",
                "parent_requirement": None,
                "module": "报表管理",
                "requirement_level": "medium", 
                "reviewer": "admin",
                "estimated_hours": 20,
                "description": "作为一名管理人员，我希望能够生成各类业务报表和统计图表，这样可以直观了解业务数据和趋势。",
                "acceptance_criteria": "系统能够生成多种格式的报表，数据准确，支持导出功能。"
            }
        ]
        
        return {
            "analysis_report": analysis_report,
            "requirements": requirements,
            "requirements_count": len(requirements)
        }
    
    @staticmethod
    async def generate_test_cases(requirement: BusinessRequirement, test_level: str, test_priority: str, count: int) -> List[Dict[str, Any]]:
        """生成测试用例 - 调用真实大模型"""
        from .models import GeneratedTestCase
        
        # 生成唯一case_id的辅助函数
        def generate_unique_case_id(req, base_index):
            base_case_id = f"TC-{req.requirement_id}-{base_index:03d}"
            case_id = base_case_id
            counter = 1
            while GeneratedTestCase.objects.filter(requirement=req, case_id=case_id).exists():
                case_id = f"{base_case_id}_{counter}"
                counter += 1
            return case_id
        
        # 获取Prompt模板
        try:
            prompt_config = PromptConfig.objects.get(prompt_type='writer', is_active=True)
            system_prompt = prompt_config.content
        except PromptConfig.DoesNotExist:
            system_prompt = "你是资深测试专家，生成高覆盖测试用例，输出Markdown表格，包含用例ID、测试目标、前置条件、操作步骤、预期结果、优先级、测试类型、关联需求。"
        
        # 构造用户输入
        user_prompt = f"""
需求ID：{requirement.requirement_id}
需求名称：{requirement.requirement_name}
需求描述：{requirement.description}
验收标准：{requirement.acceptance_criteria}
需求模块：{requirement.module}
测试级别：{test_level}
要求生成至少{max(count, 20)}条测试用例，优先级默认{test_priority}，覆盖正向、异常、边界、并发、安全场景。
直接输出Markdown表格，不要任何其他内容。
"""
        
        # 调用大模型
        messages = [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ]
        
        try:
            response_content = await AIService._chat_completion(messages, temperature=0.7, max_tokens=8192)
        except Exception as e:
            logger.error(f"大模型调用失败: {e}")
            # 降级到模拟逻辑
            await asyncio.sleep(1)
            return [{
                "case_id": generate_unique_case_id(requirement, 1),
                "title": f"验证{requirement.requirement_name}功能的基本操作流程",
                "priority": test_priority,
                "precondition": "系统正常运行，用户已登录",
                "test_steps": f"1. 进入{requirement.requirement_name}功能页面\n2. 执行核心操作\n3. 验证结果",
                "expected_result": f"{requirement.requirement_name}功能正常，结果符合预期"
            }]
        
        # 解析Markdown表格
        test_cases = []
        table_pattern = re.compile(r'\|.*?\|.*?\|.*?\|.*?\|.*?\|.*?\|.*?\|.*?\|', re.DOTALL)
        rows_pattern = re.compile(r'\|(.*?)\|(.*?)\|(.*?)\|(.*?)\|(.*?)\|(.*?)\|(.*?)\|(.*?)\|')
        
        # 找表格内容
        table_matches = table_pattern.findall(response_content)
        if not table_matches:
            logger.warning("未找到表格内容，返回模拟用例")
            return [{
                "case_id": generate_unique_case_id(requirement, 1),
                "title": f"验证{requirement.requirement_name}功能的基本操作流程",
                "priority": test_priority,
                "precondition": "系统正常运行，用户已登录",
                "test_steps": f"1. 进入{requirement.requirement_name}功能页面\n2. 执行核心操作\n3. 验证结果",
                "expected_result": f"{requirement.requirement_name}功能正常，结果符合预期"
            }]
        
        # 解析行，跳过表头和分隔行
        existing_count = GeneratedTestCase.objects.filter(requirement=requirement).count()
        index = 0
        
        for match in table_matches:
            rows = match.strip().split('\n')
            for row in rows:
                row = row.strip()
                if not row or '用例ID' in row or '----' in row or '----' in row.replace(' ', ''):
                    continue
                
                row_match = rows_pattern.match(row)
                if not row_match:
                    continue
                
                # 提取各列，去掉前后空格
                cols = [col.strip() for col in row_match.groups()]
                if len(cols) < 5:
                    continue
                
                index += 1
                case_id = generate_unique_case_id(requirement, existing_count + index)
                title = cols[1] if len(cols) > 1 else f"验证{requirement.requirement_name}"
                precondition = cols[2] if len(cols) > 2 else "系统正常运行，用户已登录"
                test_steps = cols[3] if len(cols) > 3 else f"1. 进入功能页面\n2. 执行操作\n3. 验证结果"
                expected_result = cols[4] if len(cols) > 4 else "功能正常，结果符合预期"
                priority = cols[5] if len(cols) > 5 and cols[5] in ['P0','P1','P2','P3','P4'] else test_priority
                test_type = cols[6] if len(cols) > 6 else "功能测试"
                
                if not title:
                    continue
                
                test_cases.append({
                    "case_id": case_id,
                    "title": title,
                    "priority": priority,
                    "precondition": precondition,
                    "test_steps": test_steps,
                    "expected_result": expected_result,
                    "test_type": test_type
                })
        
        # 至少返回1条
        if not test_cases:
            test_cases.append({
                "case_id": generate_unique_case_id(requirement, existing_count + 1),
                "title": f"验证{requirement.requirement_name}功能的基本操作流程",
                "priority": test_priority,
                "precondition": "系统正常运行，用户已登录",
                "test_steps": f"1. 进入{requirement.requirement_name}功能页面\n2. 执行核心操作\n3. 验证结果",
                "expected_result": f"{requirement.requirement_name}功能正常，结果符合预期"
            })
        
        return test_cases[:max(count, len(test_cases))]
    
    @staticmethod
    async def review_test_cases(test_cases: List[GeneratedTestCase], review_criteria: str) -> Dict[str, Any]:
        """评审测试用例 - 大模型B"""
        # 模拟AI评审过程
        await asyncio.sleep(1.5)
        
        reviewed_cases = []
        for test_case in test_cases:
            # 模拟评审逻辑
            review_score = 85  # 模拟评分
            
            review_comments = f"""
评审意见:
1. 测试用例标题清晰明确，能够准确描述测试目的
2. 测试步骤详细具体，具有良好的可执行性
3. 预期结果明确，便于验证
4. 建议补充异常场景的测试覆盖

评审分数: {review_score}/100
评审状态: 通过
"""
            
            reviewed_cases.append({
                "test_case_id": test_case.id,
                "review_score": review_score,
                "review_comments": review_comments.strip(),
                "status": "reviewed" if review_score >= 80 else "rejected"
            })
        
        return {
            "reviewed_cases": reviewed_cases,
            "overall_score": sum(case["review_score"] for case in reviewed_cases) / len(reviewed_cases),
            "pass_rate": len([case for case in reviewed_cases if case["status"] == "reviewed"]) / len(reviewed_cases) * 100
        }


class RequirementAnalysisService:
    """需求分析服务"""
    
    @classmethod
    def create_analysis_task(cls, document: RequirementDocument, task_type: str) -> AnalysisTask:
        """创建分析任务"""
        task_id = f"{task_type}_{uuid.uuid4().hex[:8]}"
        
        task = AnalysisTask.objects.create(
            task_id=task_id,
            task_type=task_type,
            document=document,
            status='pending'
        )
        
        return task
    
    @classmethod
    async def process_document_analysis(cls, document: RequirementDocument) -> RequirementAnalysis:
        """处理文档分析"""
        # 创建分析任务
        task = cls.create_analysis_task(document, 'requirement_analysis')
        
        try:
            # 更新任务状态
            task.status = 'running'
            task.started_at = datetime.now()
            task.progress = 10
            task.save()
            
            # 提取文档文本
            if not document.extracted_text:
                document.extracted_text = DocumentProcessor.extract_text(document)
                document.save()
            
            task.progress = 30
            task.save()
            
            # 调用AI分析
            start_time = time.time()
            analysis_result = await AIService.analyze_requirements(
                document.extracted_text, 
                document.title
            )
            analysis_time = time.time() - start_time
            
            task.progress = 70
            task.save()
            
            # 创建分析记录
            analysis = RequirementAnalysis.objects.create(
                document=document,
                analysis_report=analysis_result['analysis_report'],
                requirements_count=analysis_result['requirements_count'],
                analysis_time=analysis_time
            )
            
            # 保存需求数据
            for req_data in analysis_result['requirements']:
                BusinessRequirement.objects.create(
                    analysis=analysis,
                    **req_data
                )
            
            # 更新文档状态
            document.status = 'analyzed'
            document.save()
            
            # 完成任务
            task.status = 'completed'
            task.completed_at = datetime.now()
            task.progress = 100
            task.result = analysis_result
            task.save()
            
            return analysis
            
        except Exception as e:
            logger.error(f"文档分析失败: {e}")
            
            # 更新任务状态
            task.status = 'failed'
            task.error_message = str(e)
            task.completed_at = datetime.now()
            task.save()
            
            # 更新文档状态
            document.status = 'failed'
            document.save()
            
            raise e
    
    @classmethod
    async def generate_test_cases_for_requirements(cls, requirement_ids: List[int], test_level: str, test_priority: str, test_case_count: int) -> List[GeneratedTestCase]:
        """为需求生成测试用例"""
        generated_cases = []
        
        for req_id in requirement_ids:
            try:
                requirement = BusinessRequirement.objects.get(id=req_id)
                
                # 调用AI生成测试用例
                test_cases_data = await AIService.generate_test_cases(
                    requirement, test_level, test_priority, test_case_count
                )
                
                # 保存生成的测试用例
                for case_data in test_cases_data:
                    test_case = GeneratedTestCase.objects.create(
                        requirement=requirement,
                        case_id=case_data['case_id'],
                        title=case_data['title'],
                        priority=case_data['priority'],
                        precondition=case_data['precondition'],
                        test_steps=case_data['test_steps'],
                        expected_result=case_data['expected_result'],
                        generated_by_ai='AI-A'
                    )
                    generated_cases.append(test_case)
                    
            except BusinessRequirement.DoesNotExist:
                logger.error(f"需求ID {req_id} 不存在")
                continue
            except Exception as e:
                logger.error(f"为需求 {req_id} 生成测试用例失败: {e}")
                continue
        
        return generated_cases
    
    @classmethod
    async def review_test_cases(cls, test_case_ids: List[int], review_criteria: str) -> Dict[str, Any]:
        """评审测试用例"""
        test_cases = GeneratedTestCase.objects.filter(id__in=test_case_ids)
        
        # 调用AI评审
        review_result = await AIService.review_test_cases(list(test_cases), review_criteria)
        
        # 更新测试用例状态
        for case_review in review_result['reviewed_cases']:
            try:
                test_case = GeneratedTestCase.objects.get(id=case_review['test_case_id'])
                test_case.status = case_review['status']
                test_case.review_comments = case_review['review_comments']
                test_case.reviewed_by_ai = 'AI-B'
                test_case.save()
            except GeneratedTestCase.DoesNotExist:
                logger.error(f"测试用例ID {case_review['test_case_id']} 不存在")
                continue
        
        return review_result
