from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import filters
from rest_framework.parsers import MultiPartParser, FormParser
from django.db import models, transaction
from django.db import connection, DatabaseError
from django.db.models import OuterRef, Subquery, Q
from django.utils import timezone
import json
import re
import logging
import threading
import asyncio
import os
import uuid
import time
from urllib.parse import urlparse

from ..models import (
    UiProject, LocatorStrategy, Element, TestCase, TestCaseStep,
    AICase, AICaseCategory, AIExecutionRecord, SetupTemplate, DesktopAppProfile
)
from ..serializers import (
    AICaseCategorySerializer, AICaseSerializer, AIExecutionRecordSerializer, AIExecutionRecordListSerializer,
    DesktopAppProfileSerializer,
)
from ..structured_actions import compile_task_description
from ._helpers import (
    StandardPagination, extract_step_info, STOP_SIGNALS,
    update_planned_task_status, backfill_prior_pending_tasks,
    mark_first_active_task, summarize_planned_tasks,
    append_execution_summary, is_infrastructure_failure,
    _run_with_playwright_sync, _find_element_for_type,
    _capture_element_info,
)

logger = logging.getLogger(__name__)


_COMMUNITY_URL_PATTERN = re.compile(r'https?://[^\s，。；)]+', re.IGNORECASE)
_COMMUNITY_ROUTE_PATTERN = re.compile(
    r'(?P<route>/community/[^\s\u3002\uff0c\uff1b\)\]\}]+)',
    re.IGNORECASE,
)
_NN_COMMUNITY_HINTS = (
    '\u793e\u533a', '\u7edd\u5730\u6c42\u751f', 'PUBG', 'test0915', '\u65d7\u4e0b\u793e\u533a',
)


def _is_article_write_case(case, execution_target=''):
    """Article creation is a write workflow and must not burst concurrently."""
    target = str(execution_target or getattr(case, 'execution_target', '') or 'web').lower()
    description = (
        getattr(case, 'web_task_description', '') if target == 'web'
        else getattr(case, 'desktop_task_description', '')
    ) or getattr(case, 'task_description', '')
    text = str(description or '')
    return '文章' in text and any(token in text for token in ('发布', '存草稿', '草稿箱'))


def _desktop_task_from_web_task(task_description):
    """Translate NN web/community tasks into a runnable NN client task."""
    source = str(task_description or '').strip()
    if not source:
        return ''

    def replace_url(match):
        raw_url = match.group(0)
        parsed = urlparse(raw_url)
        host = (parsed.hostname or '').casefold()
        is_nn_host = host == 'nn.com' or host.endswith('.nn.com')
        if not is_nn_host:
            return raw_url
        if parsed.path.startswith('/community/'):
            route = parsed.path + (f'?{parsed.query}' if parsed.query else '')
            return f'NN 客户端 {route}'
        if parsed.path in ('', '/'):
            return 'NN 客户端'
        return raw_url

    translated = _COMMUNITY_URL_PATTERN.sub(replace_url, source)
    if translated != source:
        return translated

    # Older cases often start at the NN community home implicitly and only
    # describe the search/navigation flow. Keep those cases executable instead
    # of fabricating an unsupported deep link.
    if any(hint.casefold() in source.casefold() for hint in _NN_COMMUNITY_HINTS):
        return '\u542f\u52a8\u6216\u5207\u6362\u5230 NN \u5ba2\u6237\u7aef\uff0c\u8fdb\u5165\u793e\u533a\u9996\u9875\u540e\u518d\u6267\u884c\u4ee5\u4e0b\u6b65\u9aa4\uff1a\n' + source
    return ''


def _web_task_from_desktop_task(task_description, web_base_url=''):
    """Translate an NN client community route into an independent browser task."""
    source = str(task_description or '').strip()
    if not source:
        return ''

    for match in _COMMUNITY_URL_PATTERN.finditer(source):
        parsed = urlparse(match.group(0))
        host = (parsed.hostname or '').casefold()
        if (host == 'nn.com' or host.endswith('.nn.com')) and parsed.path.startswith('/community/'):
            return source

    route_match = _COMMUNITY_ROUTE_PATTERN.search(source)
    if not route_match:
        return ''

    base_url = str(web_base_url or '').rstrip('/')
    parsed_base = urlparse(base_url)
    host = (parsed_base.hostname or '').casefold()
    if not base_url or not (host == 'nn.com' or host.endswith('.nn.com')):
        base_url = 'https://test3-web.nn.com'

    route = route_match.group('route').rstrip('.,;:')
    browser_url = f'{base_url}{route}'
    translated = source[:route_match.start('route')] + browser_url + source[route_match.end('route'):]
    translated = re.sub(
        r'(\u76ee\u6807\u5730\u5740\s*[:\uff1a]\s*)NN\s*(?:\u5ba2\u6237\u7aef|client)\s*',
        r'\1',
        translated,
        count=1,
        flags=re.IGNORECASE,
    )
    translated = re.sub(
        r'NN\s*(?:\u5ba2\u6237\u7aef|client)\s*',
        '\u6d4f\u89c8\u5668 ',
        translated,
        count=1,
        flags=re.IGNORECASE,
    )
    return translated


def _infer_nn_web_base_url(task_descriptions):
    """Reuse the project's NN web host so generated tasks stay in its environment."""
    for task_description in task_descriptions:
        for match in _COMMUNITY_URL_PATTERN.finditer(str(task_description or '')):
            parsed = urlparse(match.group(0))
            host = (parsed.hostname or '').casefold()
            if (host == 'nn.com' or host.endswith('.nn.com')) and parsed.scheme and parsed.netloc:
                return f'{parsed.scheme}://{parsed.netloc}'
    return 'https://test3-web.nn.com'


def _nn_cdp_user_data_dir(ai_case, profile_config, executable):
    """Keep the automation client isolated from the user's normal NN session."""
    case_config = getattr(ai_case, 'desktop_config', None) or {}
    configured_dir = str(
        case_config.get('cdp_user_data_dir')
        or os.environ.get('TESTHUB_NN_CDP_USER_DATA_DIR')
        or ''
    ).strip()
    if configured_dir:
        return configured_dir
    base_dir = str(profile_config.get('work_dir') or os.path.dirname(executable) or '').strip()
    if not base_dir:
        return ''
    return os.path.join(base_dir, '.testhub-nn-cdp-profile')


def _default_nn_cdp_config(ai_case, profile_config):
    """Make the registered NN client runnable without per-case JSON wiring."""
    profile = getattr(ai_case, 'desktop_app_profile', None)
    executable = str(getattr(profile, 'executable_path', '') or profile_config.get('executable_path') or '').strip()
    process_name = str(getattr(profile, 'process_name', '') or profile_config.get('process_name') or '').strip()
    is_nn = os.path.basename(executable).casefold() == 'nn.exe' or process_name.casefold() == 'nn.exe'
    if not is_nn or str((ai_case.desktop_config or {}).get('automation_backend') or '').strip():
        return {}
    user_data_dir = _nn_cdp_user_data_dir(ai_case, profile_config, executable)
    launch_args = [
        '--remote-debugging-address=127.0.0.1',
        '--remote-debugging-port=9222',
    ]
    if user_data_dir:
        os.makedirs(user_data_dir, exist_ok=True)
        launch_args.append(f'--user-data-dir={user_data_dir}')
    config = {
        'automation_backend': 'cdp',
        'cdp_url': 'http://127.0.0.1:9222',
        # The isolated Agent is started as the current user, never through
        # RunAs.  It has its own profile, so this neither requests UAC nor
        # takes over the operator's normal NN session.
        'cdp_auto_start': True,
        'cdp_allow_runtime_autostart': True,
        'cdp_elevated_auto_start': False,
        'cdp_silent': True,
        'cdp_executable': executable,
        'cdp_launch_args': launch_args,
        'map_web_target_route': True,
    }
    route_source = str(
        getattr(ai_case, 'desktop_task_description', '')
        or getattr(ai_case, 'task_description', '')
        or ''
    )
    route = re.search(r'(?<![\w/])(/community/\d+(?:/[\w-]+)*(?:\?[^\s]+)?)', route_source)
    if route:
        config['cdp_route'] = route.group(1).rstrip('，。；;')
    return config


def _cdp_login_credentials_from_setup_template(setup_template):
    """Extract phone-login credentials from a selected setup template.

    Desktop CDP runs in an isolated Electron profile, so it cannot assume the
    operator's ordinary NN session is available.  Keep the credentials owned
    by the existing setup template rather than duplicating them in every case
    or hard-coding a test account in the executor.
    """
    if not setup_template:
        return {}
    steps = getattr(setup_template, 'steps', None) or []
    if not isinstance(steps, list):
        return {}
    fragments = []
    for step in steps:
        if isinstance(step, dict):
            fragments.extend(str(value) for value in step.values() if value is not None)
        elif step is not None:
            fragments.append(str(step))
    source = '\n'.join(fragments)
    phone_match = re.search(r'(?:手机号|手机号码|账号)\s*[:：]?\s*(1\d{10})', source)
    if not phone_match:
        phone_match = re.search(r'(?<!\d)(1\d{10})(?!\d)', source)
    code_match = re.search(r'(?:验证码|校验码|verification\s*code)\s*[:：]?\s*([0-9]{4,8})', source, re.IGNORECASE)
    if not phone_match or not code_match:
        return {}
    return {
        'phone': phone_match.group(1),
        'verification_code': code_match.group(1),
        'template_id': getattr(setup_template, 'id', None),
        'template_name': str(getattr(setup_template, 'name', '') or ''),
    }


def _desktop_startup_watchdog_seconds(desktop_config):
    """Keep native UIA startup finite even when Windows accessibility calls block."""
    raw_timeout = (desktop_config or {}).get('startup_watchdog_seconds')
    if raw_timeout is None:
        raw_timeout = (desktop_config or {}).get('launch_timeout_seconds', 20)
        raw_timeout = float(raw_timeout) + 25
    try:
        return max(30.0, min(float(raw_timeout), 120.0))
    except (TypeError, ValueError):
        return 45.0


def _start_desktop_startup_watchdog(execution_record, desktop_config):
    """Finalize startup hangs instead of leaving a report in ``running`` forever.

    UI Automation can block inside a native Windows COM call despite pywinauto's
    per-call timeout.  The worker cannot safely be force-killed from this
    process, so a separate watchdog owns the persisted execution state.
    """
    timeout_seconds = _desktop_startup_watchdog_seconds(desktop_config)
    record_id = execution_record.id

    def watch_startup():
        time.sleep(timeout_seconds)
        try:
            connection.close()
            record = AIExecutionRecord.objects.get(pk=record_id)
            logs = str(record.logs or '')
            if record.status != 'running' or record.steps_completed or '[Desktop] Startup complete;' in logs:
                return

            STOP_SIGNALS[record_id] = True
            record.status = 'failed'
            record.end_time = timezone.now()
            if record.start_time:
                record.duration = (record.end_time - record.start_time).total_seconds()
            record.logs = (
                logs
                + f'\n[Desktop] Startup timed out after {int(timeout_seconds)} seconds while connecting to the native app. '
                  'The execution was stopped so it will not remain in running state. '
                  'Check that the target app window is responsive and visible, then run again.\n'
            )
            record.save(update_fields=['status', 'end_time', 'duration', 'logs'])
            try:
                from ..ai_notifications import send_ai_execution_notification
                send_ai_execution_notification(record)
            except Exception:
                logger.exception('Failed to send desktop startup timeout notification (record=%s)', record_id)
        except AIExecutionRecord.DoesNotExist:
            return
        except Exception:
            logger.exception('Desktop startup watchdog failed (record=%s)', record_id)
        finally:
            connection.close()

    watcher = threading.Thread(target=watch_startup, name=f'desktop-startup-watchdog-{record_id}', daemon=True)
    watcher.start()


def _desktop_execution_config(ai_case):
    """Build the normalized desktop runtime configuration for one AI case."""
    profile_config = ai_case.desktop_app_profile.to_execution_config() if ai_case.desktop_app_profile else {}
    auto_config = _default_nn_cdp_config(ai_case, profile_config)
    desktop_config = {**profile_config, **auto_config, **(ai_case.desktop_config or {})}
    executable = str(profile_config.get('executable_path') or '').strip()
    process_name = str(profile_config.get('process_name') or '').strip()
    is_nn_client = os.path.basename(executable).casefold() == 'nn.exe' or process_name.casefold() == 'nn.exe'
    if is_nn_client and str(desktop_config.get('automation_backend') or '').lower() == 'cdp':
        # Normalize legacy machine-specific paths so every NN case uses the
        # same non-elevated, isolated CDP Agent.
        agent_profile = _nn_cdp_user_data_dir(ai_case, profile_config, executable)
        if agent_profile:
            os.makedirs(agent_profile, exist_ok=True)
        desktop_config.update({
            'cdp_url': 'http://127.0.0.1:9222',
            'cdp_auto_start': True,
            'cdp_allow_runtime_autostart': True,
            'cdp_elevated_auto_start': False,
            'cdp_silent': True,
            'cdp_executable': executable,
            'cdp_launch_args': [
                '--remote-debugging-address=127.0.0.1',
                '--remote-debugging-port=9222',
                *([f'--user-data-dir={agent_profile}'] if agent_profile else []),
            ],
        })
    if str(desktop_config.get('automation_backend') or '').lower() == 'cdp':
        template_credentials = _cdp_login_credentials_from_setup_template(ai_case.setup_template)
        if template_credentials and not desktop_config.get('cdp_login_credentials'):
            desktop_config['cdp_login_credentials'] = template_credentials
    if profile_config.get('window') and 'window' not in (ai_case.desktop_config or {}):
        desktop_config['window'] = profile_config['window']
    return desktop_config


def _preflight_nn_desktop_session(ai_case):
    """Validate a shared NN Agent session before any Windows record is made.

    A rejected SMS code is an environment problem, not a testcase result.  The
    preflight is deliberately limited to NN's CDP backend and runs once per
    user request; a successful profile is then reused by all sequential cases.
    """
    desktop_config = _desktop_execution_config(ai_case)
    if str(desktop_config.get('automation_backend') or '').lower() != 'cdp':
        return
    if not desktop_config.get('cdp_login_credentials'):
        return

    from ..desktop_engine import _CDPPage, _ensure_cdp_endpoint, _ensure_cdp_nn_login

    async def ensure_session():
        cdp_url = str(desktop_config.get('cdp_url') or 'http://127.0.0.1:9222')
        await _ensure_cdp_endpoint(cdp_url, desktop_config)
        page = await _CDPPage.connect(cdp_url)
        try:
            await _ensure_cdp_nn_login(page, desktop_config)
        finally:
            await page.close()

    asyncio.run(ensure_session())


def _run_ai_case_with_target(execution_record, ai_case, should_stop, execution_target=None):
    """Dispatch an AI case to its declared execution target."""
    target = execution_target or ai_case.execution_target
    if target == 'desktop':
        from ..desktop_engine import run_desktop_case
        desktop_config = _desktop_execution_config(ai_case)
        _start_desktop_startup_watchdog(execution_record, desktop_config)
        return run_desktop_case(
            execution_record,
            desktop_config,
            should_stop=should_stop,
        )
    if target != 'web':
        raise ValueError(f'Unsupported AI case execution target: {target}')
    # The execution record is the immutable snapshot of the variant selected by
    # the caller (browser versus desktop).  Reading ai_case.task_description
    # here can silently execute a stale/other-mode description after an edit.
    task_description = str(
        getattr(execution_record, 'task_description', '')
        or getattr(ai_case, 'web_task_description', '')
        or getattr(ai_case, 'task_description', '')
    )
    return _run_with_playwright_sync(
        execution_record,
        task_description,
        step_callback=None,
        should_stop=should_stop,
        setup_template=ai_case.setup_template,
    )


def _extract_failure_reason(record):
    """Return actionable failure evidence, never recording/GIF housekeeping logs."""
    if not record:
        return ''
    if record.status == 'stopped':
        for line in reversed(str(getattr(record, 'logs', '') or '').splitlines()):
            text = str(line).strip()
            if text and ('停止' in text or '未被旧后端调度' in text):
                return text[:200]
        return '任务已停止，未执行后续步骤。'
    for index, step in enumerate(getattr(record, 'steps_completed', None) or [], 1):
        if str(step.get('status') or '').lower() not in {'failed', 'error', 'stopped'}:
            continue
        error = str(step.get('error') or step.get('message') or '').strip()
        if error:
            return f'步骤 {index}：{error[:180]}'

    ignored_markers = ('录制完成', 'gif录制已生成', '截图', '==========')
    lines = [str(line).strip() for line in str(getattr(record, 'logs', '') or '').splitlines() if str(line).strip()]
    for line in reversed(lines):
        lowered = line.lower()
        if any(marker in lowered for marker in ignored_markers):
            continue
        if any(marker in lowered for marker in ('业务动作失败', '操作超时', '执行失败', '执行出错', '错误', '异常', 'timeout', 'error', 'failed', 'exception')):
            return line[:200]

    failed_tasks = [
        task for task in (getattr(record, 'planned_tasks', None) or [])
        if str(task.get('status') or '').lower() in {'failed', 'error', 'stopped'}
    ]
    if failed_tasks:
        task = failed_tasks[0]
        description = str(task.get('display_description') or task.get('description') or '').strip()
        if description:
            return f'失败步骤：{description[:180]}'
    return '执行失败，未记录到具体错误信息。'


def _build_ai_execution_report_items(records):
    """Build report rows at their reporting granularity, not raw record granularity."""
    standalone, batches = [], {}
    for record in records:
        batch_id = getattr(record, 'batch_id', None)
        if batch_id:
            batches.setdefault(str(batch_id), []).append(record)
        else:
            standalone.append(record)

    items = []
    for record in standalone:
        items.append({
            'id': record.id, 'item_type': 'record', 'case_name': record.case_name,
            'status': record.status, 'duration': record.duration,
            'start_time': record.start_time, 'end_time': record.end_time,
            'executed_by_name': getattr(getattr(record, 'executed_by', None), 'username', '') or '',
            'project_name': getattr(getattr(record, 'project', None), 'name', '') or '',
            'gif_path': record.gif_path, 'batch_id': None, 'case_count': 1,
        })

    terminal_statuses = {'passed', 'failed', 'stopped', 'error'}
    for batch_id, batch_records in batches.items():
        statuses = [record.status for record in batch_records]
        passed = sum(status == 'passed' for status in statuses)
        failed = sum(status in {'failed', 'error'} for status in statuses)
        stopped = sum(status == 'stopped' for status in statuses)
        pending = sum(status == 'pending' for status in statuses)
        running = sum(status == 'running' for status in statuses)
        batch_status = 'running' if running else ('pending' if pending else ('failed' if failed else ('stopped' if stopped else 'passed')))
        starts = [record.batch_started_at or record.start_time for record in batch_records if record.batch_started_at or record.start_time]
        ends = [record.end_time or record.start_time for record in batch_records if record.end_time or record.start_time]
        start_time, end_time = (min(starts), max(ends)) if starts and ends else (None, None)
        project_names = sorted({getattr(getattr(record, 'project', None), 'name', '') for record in batch_records})
        executor_names = sorted({getattr(getattr(record, 'executed_by', None), 'username', '') for record in batch_records})
        items.append({
            'id': f'batch:{batch_id}', 'item_type': 'batch', 'batch_id': batch_id,
            'case_name': f'批量执行（{len(batch_records)} 条用例）', 'status': batch_status,
            'duration': (end_time - start_time).total_seconds() if start_time and end_time else None,
            'start_time': start_time, 'end_time': end_time,
            'executed_by_name': '、'.join(name for name in executor_names if name),
            'project_name': '、'.join(name for name in project_names if name),
            'case_count': len(batch_records),
            'summary': {
                'total': len(batch_records), 'passed': passed, 'failed': failed,
                'stopped': stopped, 'running': running, 'pending': pending,
                'complete': all(status in terminal_statuses for status in statuses),
            },
        })
    return sorted(items, key=lambda item: item.get('start_time'), reverse=True)


def _build_ai_batch_report(batch_id, records):
    """Build a persisted batch summary from the member execution records."""
    started_at = next((record.batch_started_at for record in records if record.batch_started_at), None)
    if not started_at:
        started_at = min((record.start_time for record in records if record.start_time), default=None)
    finished_times = [record.end_time for record in records if record.end_time]
    finished_at = max(finished_times) if finished_times else None
    total_duration = (
        max(0, (finished_at - started_at).total_seconds())
        if started_at and finished_at else None
    )
    status_counts = {
        key: sum(1 for record in records if record.status == key)
        for key in ('passed', 'failed', 'error', 'stopped', 'running', 'pending')
    }
    finished_count = len(records) - status_counts['running'] - status_counts['pending']
    passed_count = status_counts['passed']
    failed_count = status_counts['failed'] + status_counts['error']
    return {
        'batch': {
            'id': str(batch_id),
            'run_mode': next((record.batch_run_mode for record in records if record.batch_run_mode), ''),
            'started_at': started_at,
            'finished_at': finished_at,
            'duration': total_duration,
            'is_complete': finished_count == len(records),
        },
        'summary': {
            'total': len(records),
            'passed': passed_count,
            'failed': failed_count,
            'stopped': status_counts['stopped'],
            'running': status_counts['running'],
            'pending': status_counts['pending'],
            'completion_rate': round(passed_count / len(records) * 100, 2) if records else 0,
        },
        'cases': [
            {
                'id': record.id,
                'case_name': record.case_name,
                'status': record.status,
                'duration': record.duration,
                'start_time': record.start_time,
                'end_time': record.end_time,
                'failure_reason': _extract_failure_reason(record) if record.status in ('failed', 'error', 'stopped') else '',
                'step_count': len(record.steps_completed or []),
            }
            for record in records
        ],
    }


def _normalize_recorded_actions(actions):
    if not actions:
        return []

    normalized = []
    for raw_action in actions:
        action = dict(raw_action or {})
        action_type = str(action.get('action_type') or '').lower().strip()
        if not action_type:
            continue

        locator_strategy = str(action.get('locator_strategy') or 'css').strip() or 'css'
        locator_value = str(action.get('locator_value') or '').strip()
        element_name = str(action.get('element_name') or action.get('element_text') or '').strip()
        description = str(action.get('description') or '').strip()
        input_value = str(action.get('input_value') or '')
        optional_note = str(action.get('optional_note') or '').strip()

        if action_type in {'goto', 'navigate'}:
            continue
        if action_type == 'press' and input_value in {'Shift', 'Control', 'Alt', 'Meta'}:
            continue
        if action_type == 'fill' and not locator_value and not element_name:
            continue
        if action_type == 'click' and not locator_value and not element_name and not description:
            continue

        action['action_type'] = 'fill' if action_type == 'type' else action_type
        action['locator_strategy'] = locator_strategy
        action['locator_value'] = locator_value
        action['element_name'] = element_name
        action['description'] = description
        action['input_value'] = input_value
        action['optional_note'] = optional_note

        if action['action_type'] == 'fill' and normalized:
            prev = normalized[-1]
            if prev.get('action_type') == 'fill' and prev.get('locator_value') == locator_value:
                normalized[-1] = action
                continue
            if prev.get('action_type') == 'click' and prev.get('locator_value') == locator_value and not prev.get('is_optional'):
                normalized[-1] = action
                continue

        if action['action_type'] == 'click' and normalized:
            prev = normalized[-1]
            if prev.get('action_type') == 'click':
                same_locator = prev.get('locator_value') and prev.get('locator_value') == locator_value
                same_name = prev.get('element_name') and prev.get('element_name') == element_name
                if same_locator or same_name:
                    normalized[-1] = action
                    continue

        normalized.append(action)

    cleaned = []
    for action in normalized:
        if cleaned and action.get('action_type') == 'wait' and cleaned[-1].get('action_type') == 'wait':
            cleaned[-1] = action
            continue
        cleaned.append(action)

    for idx, action in enumerate(cleaned, 1):
        action['step_number'] = idx

    return cleaned


class AICaseCategoryViewSet(viewsets.ModelViewSet):
    """Manage globally shared functional categories for AI test cases."""
    queryset = AICaseCategory.objects.all()
    permission_classes = [IsAuthenticated]
    serializer_class = AICaseCategorySerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = []
    search_fields = ['name', 'description']
    ordering_fields = ['sort_order', 'name', 'created_at', 'updated_at']
    ordering = ['sort_order', 'name', 'id']

    def get_queryset(self):
        user = self.request.user
        accessible_projects = UiProject.objects.filter(
            models.Q(owner=user) | models.Q(members=user)
        ).distinct()
        return AICaseCategory.objects.filter(
            models.Q(project__isnull=True) | models.Q(project__in=accessible_projects)
        ).select_related('project').annotate(
            cases_count=models.Count('ai_cases')
        )

    def perform_create(self, serializer):
        project = serializer.validated_data.get('project')
        if project and not UiProject.objects.filter(id=project.id).filter(
            models.Q(owner=self.request.user) | models.Q(members=self.request.user)).exists():
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied('没有该项目的权限。')
        serializer.save()


class DesktopAppProfileViewSet(viewsets.ModelViewSet):
    queryset = DesktopAppProfile.objects.all()
    permission_classes = [IsAuthenticated]
    serializer_class = DesktopAppProfileSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['is_active', 'backend']
    search_fields = ['name', 'process_name', 'executable_path', 'window_title_regex']
    ordering = ['name', 'id']

    def get_queryset(self):
        return DesktopAppProfile.objects.filter(created_by=self.request.user)

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)


class AICaseViewSet(viewsets.ModelViewSet):
    queryset = AICase.objects.all()
    permission_classes = [IsAuthenticated]
    serializer_class = AICaseSerializer
    filter_backends = [DjangoFilterBackend, filters.SearchFilter, filters.OrderingFilter]
    filterset_fields = ['project', 'category']
    search_fields = ['name', 'description', 'task_description']
    ordering = ['-created_at']

    def get_queryset(self):
        user = self.request.user
        accessible_projects = UiProject.objects.filter(
            models.Q(owner=user) | models.Q(members=user)
        ).distinct()
        latest_execution = AIExecutionRecord.objects.filter(
            ai_case_id=OuterRef('pk')
        ).order_by('-start_time')
        return AICase.objects.filter(
            models.Q(project__in=accessible_projects) | models.Q(project__isnull=True)
        ).select_related('project', 'created_by', 'setup_template', 'category', 'desktop_app_profile').annotate(
            latest_execution_id=Subquery(latest_execution.values('id')[:1]),
            latest_execution_status=Subquery(latest_execution.values('status')[:1]),
            latest_execution_time=Subquery(latest_execution.values('start_time')[:1]),
            latest_execution_logs=Subquery(latest_execution.values('logs')[:1]),
        ).distinct()

    def perform_create(self, serializer):
        project = serializer.validated_data.get('project')
        if project and not UiProject.objects.filter(id=project.id).filter(
            models.Q(owner=self.request.user) | models.Q(members=self.request.user)
        ).exists():
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied('没有该项目的权限。')
        serializer.save(created_by=self.request.user)

    def perform_update(self, serializer):
        project = serializer.validated_data.get('project', serializer.instance.project)
        if project and not UiProject.objects.filter(id=project.id).filter(
            models.Q(owner=self.request.user) | models.Q(members=self.request.user)
        ).exists():
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied('没有该项目的权限。')
        serializer.save()

    def perform_destroy(self, instance):
        instance.delete()

    @action(detail=True, methods=['post'], url_path='assertion-preview')
    def assertion_preview(self, request, pk=None):
        """Preflight the assertion plan before a case is saved or executed."""
        ai_case = self.get_object()
        task_description = str(request.data.get('task_description') or ai_case.task_description or '')
        steps = compile_task_description(task_description)
        assertions = []
        for index, step in enumerate(steps, 1):
            if step.get('action_type') != 'assert':
                continue
            contract = step.get('assertion_contract') or {}
            review_required = contract.get('kind') == 'review_required'
            assertions.append({
                'step_number': index,
                'description': step.get('display_description') or step.get('description') or '',
                'contract': contract,
                # This endpoint only compiles natural language. It does not
                # navigate, log in, or inspect the runtime DOM.
                'status': 'review_required' if review_required else 'compiled',
                'runtime_verified': False,
            })
        return Response({
            'case_id': ai_case.id,
            'assertions': assertions,
            'summary': {
                'total': len(assertions),
                'compiled': sum(item['status'] == 'compiled' for item in assertions),
                'review_required': sum(item['status'] == 'review_required' for item in assertions),
            },
            'scope': 'static_compile_only',
        })

    @staticmethod
    def _extract_markdown_cases(content):
        """Split Markdown imports by explicit case headings, not section headings."""
        case_heading_pattern = re.compile(
            r'^#{1,6}\s+(?=(?:[A-Za-z][A-Za-z0-9_]*-\d+|(?:测试)?用例\s*\d+)).+$',
            re.MULTILINE,
        )
        matches = list(case_heading_pattern.finditer(content))
        if matches:
            return [
                content[match.start():matches[index + 1].start() if index + 1 < len(matches) else len(content)].strip()
                for index, match in enumerate(matches)
            ]
        return [
            block.strip()
            for block in re.split(r'\n(?=#{1,6}\s+)', content)
            if block.strip()
        ]

    @action(detail=False, methods=['post'], url_path='import-cases', parser_classes=[MultiPartParser, FormParser])
    def import_cases(self, request):
        project_id = request.data.get('project_id')
        category_id = request.data.get('category_id')
        file_obj = request.FILES.get('file')
        if not project_id:
            return Response({'error': '请选择所属项目'}, status=status.HTTP_400_BAD_REQUEST)
        if not file_obj:
            return Response({'error': '请上传文件'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            project = UiProject.objects.get(id=project_id)
        except UiProject.DoesNotExist:
            return Response({'error': '项目不存在'}, status=status.HTTP_404_NOT_FOUND)
        if not (project.owner_id == request.user.id or project.members.filter(id=request.user.id).exists()):
            return Response({'error': '无项目权限'}, status=status.HTTP_403_FORBIDDEN)
        category = None
        if category_id not in (None, ''):
            if not (isinstance(category_id, int) or (isinstance(category_id, str) and category_id.isdigit())):
                return Response({'error': 'category_id 参数格式错误'}, status=status.HTTP_400_BAD_REQUEST)
            category = AICaseCategory.objects.filter(id=int(category_id), project__isnull=True).first()
            if not category:
                return Response({'error': '功能分类不存在'}, status=status.HTTP_400_BAD_REQUEST)
        filename = file_obj.name.lower()
        cases = []
        try:
            if filename.endswith(('.md', '.markdown', '.txt')):
                content = file_obj.read().decode('utf-8-sig', errors='replace')
                blocks = self._extract_markdown_cases(content)
                for idx, block in enumerate(blocks, 1):
                    first = block.splitlines()[0].strip('# ').strip() if block.splitlines() else f'AI用例{idx}'
                    title = first[:200] or f'AI用例{idx}'
                    cases.append({'name': title, 'task_description': block})
            elif filename.endswith(('.xlsx', '.xlsm')):
                from openpyxl import load_workbook
                wb = load_workbook(file_obj, read_only=True, data_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if rows:
                    headers = [str(x or '').strip() for x in rows[0]]
                    for idx, row in enumerate(rows[1:], 1):
                        data = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
                        name = str(data.get('用例名称') or data.get('名称') or data.get('name') or f'AI导入用例{idx}').strip()
                        task = str(data.get('任务描述') or data.get('描述') or data.get('task_description') or '').strip()
                        if task:
                            cases.append({'name': name[:200], 'task_description': task})
            else:
                return Response({'error': '仅支持 md/markdown/txt/xlsx/xlsm 文件'}, status=status.HTTP_400_BAD_REQUEST)
            created = []
            with transaction.atomic():
                for item in cases:
                    case = AICase.objects.create(
                        project=project,
                        name=item['name'],
                        task_description=item['task_description'],
                        description='批量导入',
                        category=category,
                        created_by=request.user
                    )
                    created.append({'id': case.id, 'name': case.name})
            return Response({'message': f'成功导入 {len(created)} 条AI用例', 'created_count': len(created), 'cases': created})
        except Exception as e:
            logger.error(f'批量导入AI用例失败: {e}', exc_info=True)
            return Response({'error': f'导入失败: {str(e)}'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=False, methods=['post'], url_path='batch-enable-dual-mode')
    def batch_enable_dual_mode(self, request):
        raw_ids = request.data.get('ids') or []
        if not isinstance(raw_ids, list) or not raw_ids:
            return Response({'error': '请选择至少一条 AI 用例。'}, status=status.HTTP_400_BAD_REQUEST)
        case_ids = [
            int(case_id) for case_id in raw_ids
            if isinstance(case_id, int) or (isinstance(case_id, str) and case_id.isdigit())
        ]
        if not case_ids:
            return Response({'error': '用例 ID 格式错误。'}, status=status.HTTP_400_BAD_REQUEST)
        selected_cases = list(self.get_queryset().filter(id__in=case_ids).select_related('desktop_app_profile'))
        selected_by_id = {case.id: case for case in selected_cases}
        missing_ids = sorted(set(case_ids) - set(selected_by_id))
        browser_cases = [case for case in selected_cases if case.execution_target == 'web']
        already_dual = [
            {'id': case.id, 'name': case.name, 'reason': '已存在 Windows 版。'}
            for case in selected_cases if case.execution_target == 'dual'
        ]
        already_desktop = [
            {'id': case.id, 'name': case.name, 'reason': '该用例当前为仅 Windows 模式。'}
            for case in selected_cases if case.execution_target == 'desktop'
        ]
        nn_profile = DesktopAppProfile.objects.filter(
            created_by=request.user,
            is_active=True,
            process_name__iexact='NN.exe',
        ).first()
        if browser_cases and not nn_profile:
            return Response({'error': 'Register an active NN client in Manage Desktop Apps before generating Windows variants.'}, status=status.HTTP_400_BAD_REQUEST)

        converted = []
        skipped = [
            {'id': case_id, 'name': '', 'reason': '用例不存在或没有访问权限。'}
            for case_id in missing_ids
        ] + already_dual + already_desktop
        with transaction.atomic():
            for case in browser_cases:
                desktop_task = _desktop_task_from_web_task(case.task_description)
                if not desktop_task:
                    skipped.append({'id': case.id, 'name': case.name, 'reason': 'No NN community route was found in the browser task.'})
                    continue
                case.web_task_description = case.task_description
                case.desktop_task_description = desktop_task
                case.execution_target = 'dual'
                case.desktop_app_profile = nn_profile
                case.save(update_fields=[
                    'web_task_description', 'desktop_task_description', 'execution_target',
                    'desktop_app_profile', 'updated_at',
                ])
                converted.append({'id': case.id, 'name': case.name})
        message = (
            f'已生成 {len(converted)} 条 Windows 版用例。'
            if converted else '所选用例无需生成 Windows 版。'
        )
        return Response({
            'message': message,
            'converted_count': len(converted),
            'skipped_count': len(skipped),
            'converted': converted,
            'skipped': skipped,
        })

    @action(detail=False, methods=['post'], url_path='batch-enable-browser-mode')
    def batch_enable_browser_mode(self, request):
        """Generate a browser variant for desktop-only NN community cases."""
        raw_ids = request.data.get('ids') or []
        if not isinstance(raw_ids, list) or not raw_ids:
            return Response({'error': 'Select at least one AI case.'}, status=status.HTTP_400_BAD_REQUEST)
        case_ids = [
            int(case_id) for case_id in raw_ids
            if isinstance(case_id, int) or (isinstance(case_id, str) and case_id.isdigit())
        ]
        if not case_ids:
            return Response({'error': 'Invalid AI case IDs.'}, status=status.HTTP_400_BAD_REQUEST)

        selected_cases = list(self.get_queryset().filter(id__in=case_ids).select_related('desktop_app_profile'))
        selected_by_id = {case.id: case for case in selected_cases}
        missing_ids = sorted(set(case_ids) - set(selected_by_id))
        convertible_cases = [
            case for case in selected_cases
            if case.execution_target == 'desktop'
            or (case.execution_target == 'dual' and not str(case.web_task_description or '').strip())
        ]
        skipped = [
            {'id': case_id, 'name': '', 'reason': 'Case does not exist or is not accessible.'}
            for case_id in missing_ids
        ]
        for case in selected_cases:
            if case not in convertible_cases:
                skipped.append({
                    'id': case.id,
                    'name': case.name,
                    'reason': 'A browser version already exists.',
                })

        base_url_by_project = {}
        for case in convertible_cases:
            project_key = case.project_id
            if project_key not in base_url_by_project:
                project_tasks = self.get_queryset().filter(project_id=project_key).values_list(
                    'web_task_description', 'task_description'
                )
                task_descriptions = [value for row in project_tasks for value in row]
                base_url_by_project[project_key] = _infer_nn_web_base_url(task_descriptions)

        converted = []
        with transaction.atomic():
            for case in convertible_cases:
                desktop_task = str(case.desktop_task_description or case.task_description or '').strip()
                browser_task = _web_task_from_desktop_task(
                    desktop_task,
                    base_url_by_project[case.project_id],
                )
                if not browser_task:
                    skipped.append({
                        'id': case.id,
                        'name': case.name,
                        'reason': 'No NN community route was found in the Windows task.',
                    })
                    continue
                case.desktop_task_description = desktop_task
                case.web_task_description = browser_task
                case.task_description = browser_task
                case.execution_target = 'dual'
                case.save(update_fields=[
                    'task_description', 'web_task_description', 'desktop_task_description',
                    'execution_target', 'updated_at',
                ])
                converted.append({'id': case.id, 'name': case.name})

        message = (
            f'Generated browser variants for {len(converted)} case(s).'
            if converted else 'The selected cases do not need browser variants.'
        )
        return Response({
            'message': message,
            'converted_count': len(converted),
            'skipped_count': len(skipped),
            'converted': converted,
            'skipped': skipped,
        })

    @action(detail=False, methods=['post'], url_path='batch-delete')
    def batch_delete(self, request):
        ids = request.data.get('ids', [])
        if not isinstance(ids, list) or not ids:
            return Response({'error': '请选择要删除的 AI 用例'}, status=status.HTTP_400_BAD_REQUEST)

        case_ids = [case_id for case_id in ids if isinstance(case_id, int) or (isinstance(case_id, str) and case_id.isdigit())]
        if not case_ids:
            return Response({'error': 'ids 参数格式错误'}, status=status.HTTP_400_BAD_REQUEST)

        cases_to_delete = self.get_queryset().filter(id__in=case_ids)
        deletable_ids = list(cases_to_delete.values_list('id', flat=True))
        if not deletable_ids:
            return Response({'error': '未找到可删除的 AI 用例或没有项目权限'}, status=status.HTTP_404_NOT_FOUND)

        AICase.objects.filter(id__in=deletable_ids).delete()
        return Response({
            'message': f'成功删除 {len(deletable_ids)} 条 AI 用例',
            'deleted_count': len(deletable_ids),
        })

    @action(detail=False, methods=['post'], url_path='batch-category')
    def batch_category(self, request):
        ids = request.data.get('ids', [])
        category_id = request.data.get('category_id')
        if not isinstance(ids, list) or not ids:
            return Response({'error': '请选择要配置功能分类的 AI 用例。'}, status=status.HTTP_400_BAD_REQUEST)

        case_ids = [
            int(case_id) for case_id in ids
            if isinstance(case_id, int) or (isinstance(case_id, str) and case_id.isdigit())
        ]
        if not case_ids:
            return Response({'error': 'ids 参数格式错误。'}, status=status.HTTP_400_BAD_REQUEST)

        cases = list(self.get_queryset().filter(id__in=case_ids).values('id', 'project_id'))
        if len(cases) != len(set(case_ids)):
            return Response({'error': '部分 AI 用例不存在或没有项目权限。'}, status=status.HTTP_403_FORBIDDEN)

        category = None
        if category_id not in (None, ''):
            if not (isinstance(category_id, int) or (isinstance(category_id, str) and category_id.isdigit())):
                return Response({'error': 'category_id 参数格式错误。'}, status=status.HTTP_400_BAD_REQUEST)
            accessible_projects = UiProject.objects.filter(
                models.Q(owner=request.user) | models.Q(members=request.user)
            ).distinct()
            category = AICaseCategory.objects.filter(id=int(category_id)).filter(
                models.Q(project__isnull=True) | models.Q(project__in=accessible_projects)
            ).first()
            if not category:
                return Response({'error': '功能分类不存在或没有项目权限。'}, status=status.HTTP_404_NOT_FOUND)

        cases_to_update = AICase.objects.filter(id__in=[case['id'] for case in cases])
        if category:
            updated_count = cases_to_update.update(category=category)
        else:
            updated_count = cases_to_update.update(category=None)
        return Response({
            'message': (
                f'已将 {updated_count} 条 AI 用例归入“{category.name}”。'
                if category else f'已清空 {updated_count} 条 AI 用例的功能分类。'
            ),
            'updated_count': updated_count,
            'category_id': category.id if category else None,
            'category_name': category.name if category else None,
        })

    @action(detail=False, methods=['post'], url_path='batch-setup-template')
    def batch_setup_template(self, request):
        ids = request.data.get('ids', [])
        template_id = request.data.get('setup_template_id')
        if not isinstance(ids, list) or not ids:
            return Response({'error': '请选择要配置预置模板的 AI 用例'}, status=status.HTTP_400_BAD_REQUEST)

        case_ids = [
            int(case_id) for case_id in ids
            if isinstance(case_id, int) or (isinstance(case_id, str) and case_id.isdigit())
        ]
        if not case_ids:
            return Response({'error': 'ids 参数格式错误'}, status=status.HTTP_400_BAD_REQUEST)

        cases = list(self.get_queryset().filter(id__in=case_ids).values('id', 'project_id'))
        if len(cases) != len(set(case_ids)):
            return Response({'error': '部分 AI 用例不存在或没有项目权限'}, status=status.HTTP_403_FORBIDDEN)

        template = None
        if template_id not in (None, ''):
            if not (isinstance(template_id, int) or (isinstance(template_id, str) and template_id.isdigit())):
                return Response({'error': 'setup_template_id 参数格式错误'}, status=status.HTTP_400_BAD_REQUEST)
            accessible_projects = UiProject.objects.filter(
                models.Q(owner=request.user) | models.Q(members=request.user)
            ).distinct()
            template = SetupTemplate.objects.filter(
                id=int(template_id),
                project__in=accessible_projects,
            ).select_related('project').first()
            if not template:
                return Response({'error': '预置模板不存在或没有项目权限'}, status=status.HTTP_404_NOT_FOUND)

        updated_count = AICase.objects.filter(id__in=[case['id'] for case in cases]).update(setup_template=template)
        return Response({
            'message': (
                f'已为 {updated_count} 条 AI 用例配置预置模板“{template.name}”'
                if template else f'已清空 {updated_count} 条 AI 用例的预置模板'
            ),
            'updated_count': updated_count,
            'setup_template_id': template.id if template else None,
            'setup_template_name': template.name if template else None,
        })

    @action(detail=True, methods=['post'])
    def run(self, request, pk=None):
        ai_case = self.get_object()
        requested_target = str(request.data.get('execution_target') or '').strip().lower()
        if ai_case.execution_target == 'dual':
            if requested_target not in {'web', 'desktop'}:
                return Response(
                    {'error': 'Dual-mode cases require execution_target=web or execution_target=desktop.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            execution_target = requested_target
            task_description = (
                ai_case.web_task_description
                if execution_target == 'web'
                else ai_case.desktop_task_description
            )
        else:
            if requested_target and requested_target != ai_case.execution_target:
                return Response(
                    {'error': f'This case only supports execution_target={ai_case.execution_target}.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            execution_target = ai_case.execution_target
            task_description = ai_case.task_description
        if execution_target == 'desktop':
            try:
                _preflight_nn_desktop_session(ai_case)
            except Exception as exc:
                return Response(
                    {'error': f'Windows Agent 登录态不可用：{exc}'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
        execution_record = AIExecutionRecord.objects.create(
            project=ai_case.project,
            ai_case=ai_case,
            case_name=ai_case.name,
            task_description=task_description,
            execution_target=execution_target,
            status='running',
            executed_by=request.user,
            logs="准备执行任务...\n"
        )

        def run_task():
            STOP_SIGNALS[execution_record.id] = False
            try:
                connection.close()
            except Exception:
                pass
            os.environ['DJANGO_ALLOW_ASYNC_UNSAFE'] = 'true'

            def safe_save(record, update_fields=None, max_retries=3):
                for attempt in range(max_retries):
                    try:
                        record.save(update_fields=update_fields)
                        return True
                    except (DatabaseError, Exception) as e:
                        error_str = str(e)
                        if '2006' in error_str or 'MySQL server has gone away' in error_str or '0' == error_str:
                            if attempt < max_retries - 1:
                                logger.warning(f"数据库连接失败 (尝试 {attempt + 1}/{max_retries}): {e}")
                                try:
                                    connection.close()
                                except Exception:
                                    pass
                                import time
                                time.sleep(0.5)
                                continue
                            else:
                                logger.error(f"数据库保存失败，已达最大重试次数: {e}")
                                raise
                        else:
                            logger.error(f"数据库保存失败: {e}")
                            raise
                return False

            try:
                def should_stop():
                    return STOP_SIGNALS.get(execution_record.id, False)

                _run_ai_case_with_target(execution_record, ai_case, should_stop, execution_target=execution_target)
            except Exception as e:
                error_message = str(e)
                execution_record.status = 'failed'
                execution_record.end_time = timezone.now()
                execution_record.duration = (execution_record.end_time - execution_record.start_time).total_seconds()
                execution_record.logs += f"\n执行出错: {error_message}"
                try:
                    safe_save(execution_record)
                except Exception:
                    logger.error(f"保存失败状态时出错: {e}")
            finally:
                if execution_record.id in STOP_SIGNALS:
                    del STOP_SIGNALS[execution_record.id]

        thread = threading.Thread(target=run_task)
        thread.daemon = True
        thread.start()

        return Response({
            'message': 'AI 用例开始执行',
            'execution_id': execution_record.id
        })

    @action(detail=False, methods=['post'], url_path='batch-run')
    def batch_run(self, request):
        case_ids = request.data.get('case_ids', [])
        if not case_ids:
            return Response({'error': '请选择要执行的AI用例'}, status=status.HTTP_400_BAD_REQUEST)
        ai_cases = AICase.objects.filter(id__in=case_ids)
        if not ai_cases.exists():
            return Response({'error': '未找到指定的AI用例'}, status=status.HTTP_404_NOT_FOUND)
        if ai_cases.filter(execution_target='dual').exists():
            return Response(
                {'error': 'Dual-mode cases need a target selection. Run them individually and choose Browser or Windows.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if ai_cases.filter(execution_target='desktop').exists():
            return Response(
                {'error': 'Windows 桌面用例依赖前台焦点，不支持此并发批量接口；请使用“依次执行”模式。'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if any(_is_article_write_case(case, 'web') for case in ai_cases):
            return Response(
                {'error': '文章发布和草稿用例会写入站点数据，请改用“依次执行”模式，避免并发提交触发账号频控。'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        batch_id = uuid.uuid4()
        batch_started_at = timezone.now()
        records = []
        for ai_case in ai_cases:
            execution_record = AIExecutionRecord.objects.create(
                project=ai_case.project,
                ai_case=ai_case,
                case_name=ai_case.name,
                task_description=ai_case.task_description,
                execution_target=ai_case.execution_target,
                status='running',
                executed_by=request.user,
                logs="准备执行任务...\n",
                batch_id=batch_id,
                batch_started_at=batch_started_at,
                batch_run_mode='concurrent',
            )
            records.append({'id': execution_record.id, 'case_name': ai_case.name, 'case_id': ai_case.id, 'batch_id': str(batch_id)})

            def _run_one(rec=execution_record, case=ai_case):
                STOP_SIGNALS[rec.id] = False
                try:
                    connection.close()
                except Exception:
                    pass
                os.environ['DJANGO_ALLOW_ASYNC_UNSAFE'] = 'true'

                def safe_save(obj, update_fields=None, max_retries=3):
                    for attempt in range(max_retries):
                        try:
                            obj.save(update_fields=update_fields)
                            return True
                        except Exception:
                            if attempt < max_retries - 1:
                                try: connection.close()
                                except Exception: pass
                                import time; time.sleep(0.5)
                                continue
                            raise
                    return False

                try:
                    def should_stop():
                        return STOP_SIGNALS.get(rec.id, False)

                    _run_ai_case_with_target(rec, case, should_stop)
                except Exception as e:
                    logger.error(f"批量执行AI用例失败 (record={rec.id}): {e}", exc_info=True)
                    # `error` is not a persisted AIExecutionRecord status. Keep
                    # batch records compatible with the model and report UI.
                    rec.status = 'failed'
                    rec.logs += f"\n执行异常: {str(e)}"
                    rec.end_time = timezone.now()
                    rec.duration = (rec.end_time - rec.start_time).total_seconds()
                    try: safe_save(rec, update_fields=['status', 'logs', 'end_time', 'duration'])
                    except Exception: pass

            thread = threading.Thread(target=_run_one, daemon=True)
            thread.start()

        return Response({
            'message': f'已启动 {len(records)} 个AI用例执行',
            'records': records,
            'total': len(records),
            'batch_id': str(batch_id),
            'batch_started_at': batch_started_at.isoformat(),
        })

    @action(detail=False, methods=['post'], url_path='batch-record')
    def batch_record(self, request):
        case_ids = request.data.get('case_ids', [])
        run_mode = str(request.data.get('run_mode') or 'sequential').strip().lower()
        requested_target = str(request.data.get('execution_target') or '').strip().lower()
        if not case_ids:
            return Response({'error': '请选择要录制的AI用例'}, status=status.HTTP_400_BAD_REQUEST)
        if run_mode not in {'sequential', 'concurrent'}:
            return Response({'error': 'run_mode 仅支持 sequential 或 concurrent'}, status=status.HTTP_400_BAD_REQUEST)
        if requested_target and requested_target not in {'web', 'desktop'}:
            return Response({'error': 'execution_target 仅支持 web 或 desktop'}, status=status.HTTP_400_BAD_REQUEST)
        case_map = {
            case.id: case
            for case in self.get_queryset().filter(id__in=case_ids)
        }
        ordered_case_ids = []
        for raw_case_id in case_ids:
            try:
                case_id = int(raw_case_id)
            except (TypeError, ValueError):
                continue
            if case_id in case_map and case_id not in ordered_case_ids:
                ordered_case_ids.append(case_id)
        ai_cases = [case_map[case_id] for case_id in ordered_case_ids]
        if not ai_cases:
            return Response({'error': '未找到指定的AI用例'}, status=status.HTTP_404_NOT_FOUND)
        if not requested_target and any(case.execution_target == 'dual' for case in ai_cases):
            return Response(
                {'error': 'Dual-mode cases need a target selection. Run them individually and choose Browser or Windows.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        incompatible_cases = [
            case for case in ai_cases
            if requested_target and case.execution_target not in {'dual', requested_target}
        ]
        if incompatible_cases:
            names = '、'.join(case.name for case in incompatible_cases[:3])
            more = ' 等' if len(incompatible_cases) > 3 else ''
            return Response(
                {
                    'error': (
                        f'所选用例中包含仅支持另一运行端的用例：{names}{more}。'
                        '请只选择浏览器/Windows兼容的用例，或改选对应运行端。'
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )
        effective_targets = [requested_target or case.execution_target for case in ai_cases]
        if run_mode == 'concurrent' and 'desktop' in effective_targets:
            return Response(
                {'error': 'Windows 桌面用例依赖前台焦点，批量执行请选择“依次执行”模式。'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        if run_mode == 'concurrent' and any(
            _is_article_write_case(case, requested_target or case.execution_target)
            for case in ai_cases
        ):
            return Response(
                {'error': '文章发布和草稿用例会写入站点数据，请选择“依次执行”模式，避免并发提交触发账号频控。'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if 'desktop' in effective_targets:
            desktop_case = next(
                (case for case, target in zip(ai_cases, effective_targets) if target == 'desktop'),
                None,
            )
            if desktop_case is not None:
                try:
                    _preflight_nn_desktop_session(desktop_case)
                except Exception as exc:
                    return Response(
                        {'error': f'Windows Agent 登录态不可用：{exc}'},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

        batch_id = uuid.uuid4()
        batch_started_at = timezone.now()
        records = []
        run_jobs = []
        for ai_case in ai_cases:
            execution_target = requested_target or ai_case.execution_target
            task_description = (
                ai_case.web_task_description if execution_target == 'web'
                else ai_case.desktop_task_description
            ) or ai_case.task_description
            execution_record = AIExecutionRecord.objects.create(
                project=ai_case.project,
                ai_case=ai_case,
                case_name=ai_case.name,
                task_description=task_description,
                execution_target=execution_target,
                status='pending' if run_mode == 'sequential' else 'running',
                executed_by=request.user,
                logs="准备录制任务...\n",
                execution_mode='text',
                batch_id=batch_id,
                batch_started_at=batch_started_at,
                batch_run_mode=run_mode,
            )
            records.append({
                'id': execution_record.id,
                'case_name': ai_case.name,
                'case_id': ai_case.id,
                'status': execution_record.status,
                'batch_id': str(batch_id),
            })

            def _run_one(rec=execution_record, case=ai_case, target=execution_target):
                STOP_SIGNALS[rec.id] = False
                try:
                    connection.close()
                except Exception:
                    pass
                os.environ['DJANGO_ALLOW_ASYNC_UNSAFE'] = 'true'

                def safe_save(obj, update_fields=None, max_retries=3):
                    for attempt in range(max_retries):
                        try:
                            obj.save(update_fields=update_fields)
                            return True
                        except Exception:
                            if attempt < max_retries - 1:
                                try: connection.close()
                                except Exception: pass
                                import time; time.sleep(0.5)
                                continue
                            raise
                    return False

                try:
                    if rec.status != 'running':
                        rec.status = 'running'
                        rec.logs += '\n开始执行批量任务...\n'
                        safe_save(rec, update_fields=['status', 'logs'])

                    def should_stop():
                        return STOP_SIGNALS.get(rec.id, False)

                    _run_ai_case_with_target(rec, case, should_stop, execution_target=target)
                except Exception as e:
                    logger.error(f"批量录制AI用例失败 (record={rec.id}): {e}", exc_info=True)
                    rec.status = 'failed'
                    rec.logs += f"\n录制异常: {str(e)}"
                    rec.end_time = timezone.now()
                    rec.duration = (rec.end_time - rec.start_time).total_seconds()
                    try: safe_save(rec, update_fields=['status', 'logs', 'end_time', 'duration'])
                    except Exception: pass

            run_jobs.append(_run_one)
            if run_mode == 'concurrent':
                threading.Thread(target=_run_one, daemon=True).start()

        if run_mode == 'sequential':
            def _run_sequential():
                for run_job in run_jobs:
                    run_job()

            threading.Thread(target=_run_sequential, daemon=True).start()

        return Response({
            'message': f'已按{"依次执行" if run_mode == "sequential" else "并发执行"}模式启动 {len(records)} 个AI用例录制',
            'records': records,
            'total': len(records),
            'run_mode': run_mode,
            'execution_target': requested_target or '',
            'batch_id': str(batch_id),
            'batch_started_at': batch_started_at.isoformat(),
        })

    @action(detail=True, methods=['post'], url_path='save-recording')
    def save_recording(self, request, pk=None):
        execution_record = self.get_object()
        ai_case = execution_record.ai_case
        steps_data = _normalize_recorded_actions(request.data.get('steps', []))
        project_id = request.data.get('project_id')
        case_name = request.data.get('name', f'{ai_case.name} - 录制用例' if ai_case else '录制用例')
        if not steps_data:
            return Response({'error': '未提供录制步骤'}, status=status.HTTP_400_BAD_REQUEST)
        project = None
        if project_id:
            try:
                project = UiProject.objects.get(id=project_id)
            except UiProject.DoesNotExist:
                pass
        max_step = TestCaseStep.objects.filter(
            test_case__project=project
        ).order_by('-step_number').values_list('step_number', flat=True).first() or 0
        test_case, created = TestCase.objects.get_or_create(
            name=case_name,
            defaults={
                'project': project,
                'description': f'从AI录制用例 "{ai_case.name}" 自动生成' if ai_case else '从AI录制自动生成',
                'created_by': request.user,
                'status': 'draft',
            }
        )
        created_steps = []
        for idx, step_data in enumerate(steps_data):
            step_number = step_data.get('step_number', idx + 1)
            action_type = step_data.get('action_type', 'click')
            description = step_data.get('description', '')
            input_value = step_data.get('input_value', '')
            locator_value = step_data.get('locator_value', '')
            locator_strategy = step_data.get('locator_strategy', 'css')
            element_name = step_data.get('element_name', f'AI录制_步骤{step_number}')
            is_optional = step_data.get('is_optional', False)
            optional_note = step_data.get('optional_note', '')
            element = None
            if locator_value:
                from ..models import LocatorStrategy
                strategy_map = {s.name: s for s in LocatorStrategy}
                strategy_enum = strategy_map.get(locator_strategy, LocatorStrategy.CSS)
                from ..models import Element as TestCaseElement
                element = TestCaseElement.objects.create(
                    name=element_name,
                    locator_strategy=strategy_enum,
                    locator_value=locator_value,
                    description=description,
                )
            step = TestCaseStep.objects.create(
                test_case=test_case,
                step_number=step_number,
                action_type=action_type,
                description=description,
                input_value=input_value,
                element=element,
                is_optional=is_optional,
                optional_note=optional_note,
                status='pending',
            )
            created_steps.append({
                'id': step.id,
                'step_number': step_number,
                'action_type': action_type,
                'description': description,
            })
        return Response({
            'message': f'成功保存 {len(created_steps)} 个步骤',
            'test_case_id': test_case.id,
            'test_case_name': test_case.name,
            'steps': created_steps,
            'created': created,
        })

    def _process_gif_recording(self, execution_record, history):
        try:
            from django.conf import settings
            from datetime import datetime
            default_gif_path = os.path.join(os.getcwd(), 'agent_history.gif')
            if os.path.exists(default_gif_path):
                import shutil
                gif_dir = os.path.join(settings.MEDIA_ROOT, 'ai_recording')
                os.makedirs(gif_dir, exist_ok=True)
                timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
                safe_case_name = "".join(
                    [c if c.isalnum() or c in (' ', '_', '-') else '_' for c in execution_record.case_name])
                new_gif_filename = f"{safe_case_name}_{timestamp}.gif"
                new_gif_path = os.path.join(gif_dir, new_gif_filename)
                shutil.move(default_gif_path, new_gif_path)
                from ..gif_utils import compress_gif_in_place
                compress_gif_in_place(new_gif_path)
                relative_path = f'media/ai_recording/{new_gif_filename}'
                execution_record.gif_path = relative_path
                logger.info(f"✅ GIF recording saved to: {relative_path}")
            else:
                logger.warning(f"⚠️ GIF file not found at: {default_gif_path}")
        except Exception as e:
            logger.warning(f"⚠️ Failed to process GIF recording: {e}")

    def _auto_mark_completed_tasks(self, execution_record):
        try:
            initial_completed = 0
            initial_pending = 0
            initial_failed = 0
            initial_skipped = 0
            if execution_record.planned_tasks:
                initial_completed = len([t for t in execution_record.planned_tasks if t.get('status') == 'completed'])
                initial_pending = len([t for t in execution_record.planned_tasks if t.get('status') == 'pending'])
                initial_failed = len([t for t in execution_record.planned_tasks if t.get('status') == 'failed'])
                initial_skipped = len([t for t in execution_record.planned_tasks if t.get('status') == 'skipped'])
                logger.info(f"📊 Task status summary: {initial_completed} completed, {initial_pending} pending, {initial_failed} failed, {initial_skipped} skipped")
            logger.info("📋 Task statuses are controlled by AI agent via mark_task_complete action")
        except Exception as e:
            logger.warning(f"⚠️ Failed to summarize task statuses: {e}")


class AIExecutionRecordViewSet(viewsets.ModelViewSet):
    """AI执行记录视图集"""
    queryset = AIExecutionRecord.objects.all()
    serializer_class = AIExecutionRecordSerializer
    permission_classes = [IsAuthenticated]
    filter_backends = [DjangoFilterBackend, filters.OrderingFilter]
    filterset_fields = ['project', 'ai_case', 'status']
    ordering = ['-start_time']
    pagination_class = StandardPagination

    def get_queryset(self):
        user = self.request.user
        accessible_projects = UiProject.objects.filter(
            models.Q(owner=user) | models.Q(members=user)
        ).distinct()
        return AIExecutionRecord.objects.filter(
            models.Q(project__in=accessible_projects) | models.Q(project__isnull=True)
        ).select_related('project', 'ai_case', 'executed_by').defer('steps_completed', 'planned_tasks', 'screenshots_sequence').distinct()

    def get_serializer_class(self):
        if self.action == 'list':
            return AIExecutionRecordListSerializer
        return AIExecutionRecordSerializer

    def perform_destroy(self, instance):
        instance.delete()

    @action(detail=False, methods=['get'], url_path='batch-report')
    def batch_report(self, request):
        raw_batch_id = str(request.query_params.get('batch_id') or '').strip()
        try:
            batch_id = uuid.UUID(raw_batch_id)
        except (ValueError, TypeError, AttributeError):
            return Response({'error': 'batch_id 格式错误'}, status=status.HTTP_400_BAD_REQUEST)

        records = list(self.get_queryset().filter(batch_id=batch_id).order_by('start_time', 'id'))
        if not records:
            return Response({'error': '未找到该批量执行记录或没有权限'}, status=status.HTTP_404_NOT_FOUND)

        return Response(_build_ai_batch_report(batch_id, records))

    @action(detail=False, methods=['get'], url_path='report-list')
    def report_list(self, request):
        """Return one row per standalone execution or batch report."""
        try:
            page = max(int(request.query_params.get('page', 1)), 1)
            page_size = min(max(int(request.query_params.get('page_size', 20)), 1), 100)
        except (TypeError, ValueError):
            return Response({'error': '分页参数格式错误'}, status=status.HTTP_400_BAD_REQUEST)
        items = _build_ai_execution_report_items(self.get_queryset().order_by('-start_time'))
        start = (page - 1) * page_size
        return Response({'count': len(items), 'results': items[start:start + page_size]})

    @action(detail=False, methods=['post'])
    def batch_delete(self, request):
        try:
            ids = request.data.get('ids', [])
            batch_ids = request.data.get('batch_ids', [])
            if not ids and not batch_ids:
                return Response({'error': '请选择要删除的记录'}, status=status.HTTP_400_BAD_REQUEST)
            if not isinstance(ids, list) or not isinstance(batch_ids, list):
                return Response({'error': 'ids 和 batch_ids 参数必须为数组'}, status=status.HTTP_400_BAD_REQUEST)
            queryset = self.get_queryset()
            records_to_delete = queryset.filter(Q(id__in=ids) | Q(batch_id__in=batch_ids))
            if not records_to_delete.exists():
                return Response({'error': '未找到可删除的记录或没有权限删除'}, status=status.HTTP_404_NOT_FOUND)
            deletable_ids = list(records_to_delete.values_list('id', flat=True))
            deleted_count = AIExecutionRecord.objects.filter(id__in=deletable_ids).delete()[0]
            return Response({'message': f'成功删除 {deleted_count} 条记录', 'deleted_count': deleted_count})
        except Exception as e:
            logger.error(f"批量删除AI执行记录失败: {str(e)}", exc_info=True)
            return Response({'error': f'批量删除失败: {str(e)}'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    def _get_or_create_ai_element(self, project, meta):
        strategy, _ = LocatorStrategy.objects.get_or_create(
            name=meta.get('strategy', 'css'),
            defaults={'description': meta.get('strategy', 'css')}
        )
        element, _ = Element.objects.get_or_create(
            project=project,
            name=meta['name'],
            defaults={
                'element_type': meta.get('element_type', 'BUTTON'),
                'locator_strategy': strategy,
                'locator_value': meta['value'],
                'description': 'AI录制自动生成元素',
                'is_unique': True,
                'wait_timeout': 5,
            }
        )
        changed = False
        if element.locator_strategy_id != strategy.id:
            element.locator_strategy = strategy; changed = True
        if element.locator_value != meta['value']:
            element.locator_value = meta['value']; changed = True
        if changed:
            element.save(update_fields=['locator_strategy', 'locator_value'])
        return element

    def _infer_ai_recording_element_meta(self, action):
        desc = str(action.get('description') or action.get('text') or '').strip()
        low = desc.lower()
        if '同意并登录' in desc:
            return {'name': 'AI录制_同意并登录按钮', 'strategy': 'css', 'value': 'button:has-text("同意并登录")', 'element_type': 'BUTTON'}
        if any(k in desc for k in ['打开登录入口', '默认头像', '登录入口']):
            return {'name': 'AI录制_登录入口', 'strategy': 'css', 'value': '#aside-header-avatar', 'element_type': 'BUTTON'}
        if '手机号' in desc or '手机' in desc:
            return {'name': 'AI录制_手机号输入框', 'strategy': 'css', 'value': 'input[placeholder*="手机"]', 'element_type': 'INPUT'}
        if '验证码' in desc:
            return {'name': 'AI录制_验证码输入框', 'strategy': 'css', 'value': 'input[placeholder*="验证码"]', 'element_type': 'INPUT'}
        if ('登录按钮' in desc) or ('点击登录' in desc):
            return {'name': 'AI录制_登录按钮', 'strategy': 'xpath', 'value': "(//button[contains(normalize-space(),'登录') or contains(normalize-space(),'注册')])[last()]", 'element_type': 'BUTTON'}
        if '协议' in desc or '勾选' in desc or '同意' in desc:
            return {'name': 'AI录制_协议勾选框', 'strategy': 'xpath', 'value': "(//*[contains(concat(' ', normalize-space(@class), ' '), ' nn-checkbox ')])[1]", 'element_type': 'CHECKBOX'}
        if '关闭' in desc or '弹窗' in desc:
            return {'name': 'AI录制_关闭弹窗按钮', 'strategy': 'xpath', 'value': "(//*[contains(@class,'close') or contains(@class,'Close') or contains(@aria-label,'关闭') or contains(@aria-label,'close')])[1]", 'element_type': 'BUTTON'}
        if '断言' in desc or '验证' in desc or '检查' in desc or '确认' in desc:
            return {'name': 'AI录制_页面正文', 'strategy': 'css', 'value': 'body', 'element_type': 'CONTAINER'}
        return {'name': 'AI录制_页面正文', 'strategy': 'css', 'value': 'body', 'element_type': 'CONTAINER'}

    def _ai_recording_actions_from_execution_record(self, record):
        raw_items = []
        if record.planned_tasks:
            for t in record.planned_tasks:
                if t.get('status') in ('pending', 'in_progress'):
                    continue
                raw_items.append(str(t.get('description') or t.get('text') or '').strip())
        if not raw_items:
            def _normalize_recording_line(text):
                text = str(text or '').strip()
                text = re.sub(r'^\s*(?:步骤|step)?\s*\d+[.、:\-\s]+', '', text, flags=re.I)
                text = re.sub(r'^\s*[\-•*]\s*', '', text)
                return text.strip()

            raw_items = [
                _normalize_recording_line(x)
                for x in str(record.task_description or '').splitlines()
                if _normalize_recording_line(x)
            ]
        actions = []
        for text in raw_items:
            if not text:
                continue
            if re.search(r'https?://', text):
                continue
            action_type = 'click'
            input_value = ''
            assert_type = ''
            assert_value = ''
            is_optional = False
            if any(k in text for k in ['输入', '填写', '键入']):
                action_type = 'fill'
                m = re.search(r'(手机号|手机|验证码|密码|账号|用户名)\s*[:：]?\s*([\w@.\-+]+)', text)
                input_value = m.group(2) if m else ''
            elif '等待' in text:
                action_type = 'wait'; input_value = ''
            elif any(k in text for k in ['断言', '验证', '检查', '确认是否', '是否成功']):
                action_type = 'aiAssert'; assert_type = 'textContains'; assert_value = text; input_value = text
            if '弹窗' in text and '关闭' in text:
                is_optional = True
            actions.append({'action_type': action_type, 'description': text, 'input_value': input_value,
                            'assert_type': assert_type, 'assert_value': assert_value, 'wait_time': 1000,
                            'is_optional': is_optional, 'optional_note': 'AI录制：该弹窗不存在时跳过' if is_optional else ''})
        has_phone = any('手机' in a['description'] or '手机号' in a['description'] for a in actions)
        has_login_entry = any('登录入口' in a['description'] or '默认头像' in a['description'] or '登录' in a['description'] for a in actions)
        if has_phone and not has_login_entry:
            idx = next((i for i,a in enumerate(actions) if '手机' in a['description'] or '手机号' in a['description']), 0)
            actions.insert(idx, {'action_type':'click','description':'打开登录入口/点击默认头像','input_value':'','assert_type':'','assert_value':'','wait_time':1000,'is_optional':False,'optional_note':''})
        has_login_btn = any('点击登录按钮' in a['description'] for a in actions)
        has_assert = any(a['action_type'] == 'aiAssert' for a in actions)
        if has_phone and has_assert and not has_login_btn:
            idx = next((i for i,a in enumerate(actions) if a['action_type']=='aiAssert'), len(actions))
            actions.insert(idx, {'action_type':'click','description':'点击登录按钮','input_value':'','assert_type':'','assert_value':'','wait_time':1000,'is_optional':False,'optional_note':''})
        has_agree = any('同意并登录' in a['description'] for a in actions)
        if has_phone and has_assert and not has_agree:
            idx = next((i for i,a in enumerate(actions) if a['action_type']=='aiAssert'), len(actions))
            actions.insert(idx, {'action_type':'click','description':'同意并登录','input_value':'','assert_type':'','assert_value':'','wait_time':1000,'is_optional':True,'optional_note':'AI录制：二次协议确认弹窗不存在时跳过'})
        seen = set()
        deduped_actions = []
        for a in actions:
            key = (a['action_type'], a['description'][:30])
            if key not in seen:
                seen.add(key)
                deduped_actions.append(a)
        actions = deduped_actions
        for idx, a in enumerate(actions, 1):
            meta = self._infer_ai_recording_element_meta(a)
            a.update({'step_number': idx, 'element_name': meta['name'], 'locator_strategy': meta['strategy'], 'locator_value': meta['value'], 'element_type': meta.get('element_type','BUTTON')})
        return _normalize_recorded_actions(actions)

    @action(detail=True, methods=['get'], url_path='ai-recording-actions')
    def ai_recording_actions(self, request, pk=None):
        record = self.get_object()
        actions = self._ai_recording_actions_from_execution_record(record)
        return Response({'record_id': record.id, 'project_id': record.project_id, 'actions': actions, 'count': len(actions)})

    @action(detail=True, methods=['post'], url_path='save-ai-recording')
    def save_ai_recording(self, request, pk=None):
        record = self.get_object()
        if request.data.get('target_test_case_id'):
            return Response({'error': 'AI录制只支持新建普通UI用例，不支持追加到已有用例'}, status=status.HTTP_400_BAD_REQUEST)
        project_id = request.data.get('project_id') or record.project_id
        if not project_id:
            return Response({'error': '缺少所属项目'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            project = UiProject.objects.get(id=project_id)
        except UiProject.DoesNotExist:
            return Response({'error': '项目不存在'}, status=status.HTTP_404_NOT_FOUND)
        if not (project.owner_id == request.user.id or project.members.filter(id=request.user.id).exists()):
            return Response({'error': '无项目权限'}, status=status.HTTP_403_FORBIDDEN)
        case_name = request.data.get('case_name') or f'AI录制用例_{timezone.localtime(record.start_time).strftime("%Y/%m/%d %H:%M:%S")}'
        actions = _normalize_recorded_actions(request.data.get('actions') or self._ai_recording_actions_from_execution_record(record))
        if not actions:
            return Response({'error': '过滤后没有可保存的录制步骤'}, status=status.HTTP_400_BAD_REQUEST)
        with transaction.atomic():
            test_case = TestCase.objects.create(project=project, name=case_name, description=f'来自AI录制执行记录 #{record.id} 起始URL: {project.base_url or ""}', priority='medium', status='draft', created_by=request.user)
            for idx, action in enumerate(actions, 1):
                meta = {'name': action.get('element_name'), 'strategy': action.get('locator_strategy') or 'css', 'value': action.get('locator_value'), 'element_type': action.get('element_type','BUTTON')}
                if not meta['name'] or not meta['value']:
                    meta = self._infer_ai_recording_element_meta(action)
                element = self._get_or_create_ai_element(project, meta)
                TestCaseStep.objects.create(
                    test_case=test_case, step_number=idx, action_type=action.get('action_type','click'), element=element,
                    input_value=action.get('input_value',''), wait_time=action.get('wait_time',1000),
                    assert_type=action.get('assert_type',''), assert_value=action.get('assert_value',''),
                    is_optional=action.get('is_optional', False), optional_note=action.get('optional_note',''),
                    description=action.get('description','')
                )
        return Response({'message': '保存成功', 'test_case_id': test_case.id, 'test_case_name': test_case.name, 'steps_created': len(actions), 'save_mode': 'create', 'actions': actions})

    @action(detail=False, methods=['post'], url_path='run_adhoc')
    def run_adhoc(self, request):
        project_id = request.data.get('project_id')
        task_description = request.data.get('task_description')
        execution_mode = request.data.get('execution_mode', 'text')
        enable_gif = request.data.get('enable_gif', True)
        setup_template_id = request.data.get('setup_template_id')

        if not task_description:
            return Response({'error': '缺少任务描述参数'}, status=status.HTTP_400_BAD_REQUEST)

        project = None
        if project_id:
            try:
                project = UiProject.objects.get(id=project_id)
            except UiProject.DoesNotExist:
                return Response({'error': '项目不存在'}, status=status.HTTP_404_NOT_FOUND)

        execution_record = AIExecutionRecord.objects.create(
            project=project,
            case_name="Adhoc Task",
            task_description=task_description,
            execution_mode=execution_mode,
            status='running',
            executed_by=request.user,
            logs="正在分析任务...\n"
        )

        from asgiref.sync import sync_to_async
        from ..ai_agent import run_full_process_sync

        def run_task():
            local_task_description = task_description
            STOP_SIGNALS[execution_record.id] = False
            try:
                connection.close()
            except Exception:
                pass
            os.environ['DJANGO_ALLOW_ASYNC_UNSAFE'] = 'true'

            def safe_save(record, update_fields=None, max_retries=3):
                for attempt in range(max_retries):
                    try:
                        record.save(update_fields=update_fields)
                        return True
                    except (DatabaseError, Exception) as e:
                        error_str = str(e)
                        if '2006' in error_str or 'MySQL server has gone away' in error_str or '0' == error_str:
                            if attempt < max_retries - 1:
                                logger.warning(f"数据库连接失败 (尝试 {attempt + 1}/{max_retries}): {e}")
                                try:
                                    connection.close()
                                except Exception:
                                    pass
                                import time
                                time.sleep(0.5)
                                continue
                            else:
                                logger.error(f"数据库保存失败，已达最大重试次数: {e}")
                                raise
                        else:
                            logger.error(f"数据库保存失败: {e}")
                            raise
                return False

            try:
                async def should_stop_async():
                    if STOP_SIGNALS.get(execution_record.id, False):
                        return True
                    await sync_to_async(execution_record.refresh_from_db)()
                    return execution_record.status == 'stopped'

                def should_stop_sync():
                    if STOP_SIGNALS.get(execution_record.id, False):
                        return True
                    execution_record.refresh_from_db()
                    return execution_record.status == 'stopped'

                async def on_analysis_complete(planned_tasks):
                    execution_record.planned_tasks = planned_tasks
                    execution_record.logs += "任务分析完成，开始执行...\n"
                    await sync_to_async(safe_save)(execution_record, update_fields=['planned_tasks', 'logs'])

                async def on_step_update(step_info):
                    try:
                        if step_info.get('type') == 'log':
                            content = step_info.get('content')
                            if content:
                                content = str(content).replace(chr(13)+chr(10), chr(10)).replace(chr(13), chr(10))
                                execution_record.logs += content
                                safe_save(execution_record, update_fields=['logs'])
                            return
                        if step_info.get('type') == 'record_action':
                            action = step_info.get('action')
                            if action:
                                if not execution_record.steps_completed:
                                    execution_record.steps_completed = []
                                execution_record.steps_completed.append(action)
                                await sync_to_async(safe_save)(execution_record, update_fields=['steps_completed'])
                            return
                        task_id = step_info.get('task_id')
                        step_status = step_info.get('status')
                        logger.info(f"DEBUG: on_step_update received: task_id={task_id}, status={step_status}")
                        if task_id and step_status:
                            updated = False
                            if execution_record.planned_tasks:
                                old_status = None
                                for task in execution_record.planned_tasks:
                                    if str(task.get('id')) == str(task_id):
                                        old_status = task.get('status', 'pending')
                                        break
                                backfilled_ids = []
                                if str(step_status).strip().lower() == 'completed':
                                    backfilled_ids = backfill_prior_pending_tasks(
                                        execution_record.planned_tasks,
                                        task_id
                                    )
                                    if backfilled_ids:
                                        execution_record.logs += (
                                            f"\n[System] 已补齐遗漏标记的前序子任务: "
                                            f"{', '.join(map(str, backfilled_ids))}"
                                        )
                                updated = update_planned_task_status(
                                    execution_record.planned_tasks,
                                    task_id,
                                    step_status
                                )
                                if updated:
                                    logger.info(f"DEBUG: Updated task {task_id} from {old_status} to {step_status}")
                            if updated:
                                update_fields = ['planned_tasks']
                                if 'backfilled_ids' in locals() and backfilled_ids:
                                    update_fields.append('logs')
                                await sync_to_async(safe_save)(execution_record, update_fields=update_fields)
                            else:
                                logger.warning(
                                    f"DEBUG: Task ID {task_id} not found in planned_tasks: {execution_record.planned_tasks}")
                    except Exception as e:
                        logger.error(f"更新步骤状态失败: {e}", exc_info=True)

                _setup_template = None
                if setup_template_id:
                    try:
                        from ..models import SetupTemplate
                        from apps.core.variable_resolver import resolve_variables
                        _setup_template = SetupTemplate.objects.filter(id=setup_template_id).first()
                        if _setup_template and _setup_template.steps:
                            # 格式化前置步骤，包含动作类型和描述
                            _tpl_parts = []
                            for s in _setup_template.steps:
                                if isinstance(s, dict):
                                    action = s.get('action_type', '')
                                    desc = resolve_variables(str(s.get('description', '') or ''))
                                    if action and desc:
                                        _tpl_parts.append(f"{action}: {desc}")
                                    elif desc:
                                        _tpl_parts.append(desc)
                                    else:
                                        _tpl_parts.append(str(s))
                                else:
                                    _tpl_parts.append(resolve_variables(str(s)))
                            _tpl_text = "; ".join(_tpl_parts)
                            local_task_description = f"前置步骤：{_tpl_text}。然后执行以下任务：{local_task_description}"
                            execution_record.logs += f"\n[前置模板] {_setup_template.name}: {len(_setup_template.steps)} 步\n"
                            execution_record.logs += f"[前置步骤内容] {_tpl_text}\n"
                            safe_save(execution_record, update_fields=['logs'])
                    except Exception as e:
                        logger.warning(f"加载前置模板失败 {setup_template_id}: {e}")

                history = run_full_process_sync(
                    local_task_description,
                    analysis_callback=on_analysis_complete,
                    step_callback=on_step_update,
                    should_stop=should_stop_async,
                    execution_mode=execution_mode,
                    enable_gif=enable_gif,
                    case_name=local_task_description[:50] if local_task_description else "Adhoc Task"
                )

                if should_stop_sync():
                    execution_record.status = 'stopped'
                    execution_record.logs += "\n[System] 任务已由用户停止。"
                else:
                    failed = False
                    if execution_record.planned_tasks:
                        failed = any(t.get('status') in ('failed', 'error') for t in execution_record.planned_tasks)
                    if not failed and history:
                        try:
                            steps = history.steps if hasattr(history, 'steps') else []
                            for step in steps:
                                if getattr(step, 'status', None) == 'failed':
                                    failed = True
                                    break
                        except Exception:
                            pass
                    has_pending = False
                    if execution_record.planned_tasks:
                        has_pending = any(t.get('status') in ('pending', 'in_progress') for t in execution_record.planned_tasks)
                    if failed:
                        execution_record.status = 'failed'
                        execution_record.logs += "\n执行完成，但检测到失败步骤或任务。"
                    elif has_pending:
                        execution_record.status = 'failed'
                        pending_count = len([t for t in execution_record.planned_tasks if t.get('status') in ('pending', 'in_progress')])
                        execution_record.logs += f"\n执行完成，但有 {pending_count} 个子任务未完成（AI未标记）。"
                    else:
                        execution_record.status = 'passed'
                        execution_record.logs += "\n执行完成。"
                    if execution_record.planned_tasks:
                        total_tasks = len(execution_record.planned_tasks)
                        completed_tasks = len(
                            [t for t in execution_record.planned_tasks if t.get('status') == 'completed'])
                        pending_tasks = len([t for t in execution_record.planned_tasks if t.get('status') == 'pending'])
                        logger.info(
                            f"🏁 Task completion summary: {completed_tasks}/{total_tasks} tasks completed, {pending_tasks} pending")

                execution_record.end_time = timezone.now()
                execution_record.duration = (execution_record.end_time - execution_record.start_time).total_seconds()
                steps = []
                if history:
                    if hasattr(history, 'steps'):
                        steps = [extract_step_info(s, i) for i, s in enumerate(history.steps)]
                if steps:
                    execution_record.steps_completed = steps
                elif not execution_record.steps_completed:
                    execution_record.steps_completed = steps
                if execution_record.planned_tasks:
                    self._auto_mark_completed_tasks(execution_record)
                    execution_record.logs = append_execution_summary(
                        execution_record.logs,
                        summarize_planned_tasks(execution_record.planned_tasks)
                    )
                self._process_gif_recording(execution_record, history)
                safe_save(execution_record)

            except Exception as e:
                error_message = str(e)
                failed_task_id = None if is_infrastructure_failure(error_message) else mark_first_active_task(execution_record.planned_tasks, 'failed')
                execution_record.status = 'failed'
                execution_record.end_time = timezone.now()
                execution_record.duration = (execution_record.end_time - execution_record.start_time).total_seconds()
                if 'Execution LLM unavailable' in error_message:
                    execution_record.logs += f"\n执行出错: AI 执行模型连接失败。{error_message}"
                else:
                    execution_record.logs += f"\n执行出错: {error_message}"
                if failed_task_id is not None:
                    execution_record.logs += f"\n[System] 子任务 {failed_task_id} 已自动标记为失败。"
                execution_record.logs = append_execution_summary(
                    execution_record.logs,
                    summarize_planned_tasks(execution_record.planned_tasks)
                )
                try:
                    safe_save(execution_record)
                except Exception:
                    logger.error(f"保存失败状态时出错: {e}")
                    pass
            finally:
                try:
                    from ..ai_notifications import send_ai_execution_notification
                    send_ai_execution_notification(execution_record)
                except Exception as notify_error:
                    logger.error(f"发送 AI 测试报告通知失败: {notify_error}", exc_info=True)
                if execution_record.id in STOP_SIGNALS:
                    del STOP_SIGNALS[execution_record.id]

        thread = threading.Thread(target=run_task)
        thread.daemon = True
        thread.start()

        return Response({
            'message': 'AI 任务开始执行',
            'execution_id': execution_record.id
        })

    @action(detail=True, methods=['post'], url_path='stop')
    def stop_task(self, request, pk=None):
        try:
            execution_id = int(pk)
            if execution_id in STOP_SIGNALS:
                STOP_SIGNALS[execution_id] = True
                return Response({'message': '已发送停止信号'})
            else:
                record = self.get_object()
                if record.status == 'running':
                    record.status = 'stopped'
                    record.end_time = timezone.now()
                    record.logs += "\n[System] 任务被强制标记为停止（未在运行队列中找到）。"
                    record.save()
                    return Response({'message': '任务已标记为停止'})
                return Response({'message': '任务不在运行中'}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    def _process_gif_recording(self, execution_record, history):
        try:
            from django.conf import settings
            from datetime import datetime
            default_gif_path = os.path.join(os.getcwd(), 'agent_history.gif')
            if os.path.exists(default_gif_path):
                import shutil
                gif_dir = os.path.join(settings.MEDIA_ROOT, 'ai_recording')
                os.makedirs(gif_dir, exist_ok=True)
                timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
                safe_case_name = "".join(
                    [c if c.isalnum() or c in (' ', '_', '-') else '_' for c in execution_record.case_name])
                new_gif_filename = f"{safe_case_name}_{timestamp}.gif"
                new_gif_path = os.path.join(gif_dir, new_gif_filename)
                shutil.move(default_gif_path, new_gif_path)
                from ..gif_utils import compress_gif_in_place
                compress_gif_in_place(new_gif_path)
                relative_path = f'media/ai_recording/{new_gif_filename}'
                execution_record.gif_path = relative_path
                logger.info(f"✅ GIF recording saved to: {relative_path}")
            else:
                logger.warning(f"⚠️ GIF file not found at: {default_gif_path}")
        except Exception as e:
            logger.warning(f"⚠️ Failed to process GIF recording: {e}")

    def _auto_mark_completed_tasks(self, execution_record):
        try:
            initial_completed = 0
            initial_pending = 0
            initial_failed = 0
            initial_skipped = 0
            if execution_record.planned_tasks:
                initial_completed = len([t for t in execution_record.planned_tasks if t.get('status') == 'completed'])
                initial_pending = len([t for t in execution_record.planned_tasks if t.get('status') == 'pending'])
                initial_failed = len([t for t in execution_record.planned_tasks if t.get('status') == 'failed'])
                initial_skipped = len([t for t in execution_record.planned_tasks if t.get('status') == 'skipped'])
                logger.info(f"📊 Task status summary: {initial_completed} completed, {initial_pending} pending, {initial_failed} failed, {initial_skipped} skipped")
            logger.info("📋 Task statuses are controlled by AI agent via mark_task_complete action")
        except Exception as e:
            logger.warning(f"⚠️ Failed to summarize task statuses: {e}")

    @action(detail=True, methods=['get'], url_path='report')
    def generate_report(self, request, pk=None):
        try:
            record = self.get_object()
            report_type = request.query_params.get('report_type', 'summary')
            from ..reports import AIExecutionReportGenerator
            generator = AIExecutionReportGenerator(record)
            if report_type == 'detailed':
                report = generator.generate_detailed_report()
            elif report_type == 'performance':
                report = generator.generate_performance_report()
            else:
                report = generator.generate_summary_report()
            return Response({
                'success': True,
                'data': report,
                'report_type': report_type
            })
        except Exception as e:
            logger.error(f"生成AI执行报告失败: {e}", exc_info=True)
            return Response({
                'success': False,
                'error': str(e)
            }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=True, methods=['get'], url_path='export-pdf')
    def export_pdf(self, request, pk=None):
        from django.http import HttpResponse
        try:
            record = self.get_object()
            report_type = request.query_params.get('report_type', 'summary')
            from ..reports import AIExecutionReportGenerator
            from ..pdf_generator import AIReportPDFGenerator
            generator = AIExecutionReportGenerator(record)
            if report_type == 'detailed':
                report_data = generator.generate_detailed_report()
            elif report_type == 'performance':
                report_data = generator.generate_performance_report()
            else:
                report_data = generator.generate_summary_report()
            pdf_generator = AIReportPDFGenerator(report_data, report_type)
            pdf_buffer = pdf_generator.generate()
            from datetime import datetime
            timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
            safe_case_name = "".join([c if c.isalnum() or c in (' ', '_', '-') else '_' for c in record.case_name])
            filename = f"AI_Report_{safe_case_name}_{timestamp}.pdf"
            response = HttpResponse(
                pdf_buffer.getvalue(),
                content_type='application/pdf'
            )
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            response['Content-Length'] = len(pdf_buffer.getvalue())
            return response
        except ImportError as e:
            logger.error(f"PDF生成库未安装: {e}")
            return Response({
                'success': False,
                'error': 'PDF生成功能需要安装 reportlab 库，请运行: pip install reportlab'
            }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
        except Exception as e:
            logger.error(f"导出PDF失败: {e}", exc_info=True)
            return Response({
                'success': False,
                'error': str(e)
            }, status=status.HTTP_500_INTERNAL_SERVER_ERROR)


class AIAssertionView(APIView):
    """AI断言验证接口"""
    def post(self, request):
        from ..ai_assertion import AIAssertionService
        service = AIAssertionService()
        result = service.assert_condition(request.data)
        return Response(result)
